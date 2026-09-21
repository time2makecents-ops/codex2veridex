import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from batch_inventory import (
    ItemSelection,
    apply_price_evaluations_to_workbook,
    create_annotated_selection_image,
    create_enhancement_comparison_sheet,
    create_enhancement_variants,
    create_inventory_workbook,
    create_layered_sharpen_enhancement,
    create_review_contact_sheet,
    export_item_images,
    render_selection,
)


class BatchInventoryTests(unittest.TestCase):
    def test_lasso_creates_context_transparent_and_neutral_versions(self):
        source = Image.new("RGB", (100, 80), "blue")
        selection = ItemSelection(1, ((10, 10), (80, 10), (50, 70)))
        result = render_selection(source, selection, padding=0)
        self.assertEqual(result["context"].mode, "RGB")
        self.assertEqual(result["transparent"].mode, "RGBA")
        self.assertEqual(result["neutral"].mode, "RGB")
        self.assertEqual(result["transparent"].getpixel((0, 59))[3], 0)

    def test_workbook_embeds_items_and_review_rows(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.jpg"
            Image.new("RGB", (160, 120), "white").save(source)
            selections = [
                ItemSelection(
                    1,
                    ((10, 10), (100, 80)),
                    description="Test item",
                    review_status="Needs review",
                    unreadable_reason="Label too small",
                )
            ]
            items = export_item_images(source, selections, root / "crops")
            output = create_inventory_workbook(root / "inventory.xlsx", source, items)
            contact_sheet = create_review_contact_sheet(root / "review.png", items)
            annotated = create_annotated_selection_image(source, selections, root / "annotated.png")
            enhanced = create_layered_sharpen_enhancement(source, root / "enhanced.png")
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(workbook["Inventory"]["C2"].value, "Test item")
            self.assertEqual(workbook["Needs Review"]["C2"].value, "Label too small")
            self.assertEqual(workbook["Source & Instructions"]["A1"].value, "Source image")
            workbook.close()
            with Image.open(contact_sheet) as image:
                self.assertGreater(image.width, 0)
                self.assertGreater(image.height, 0)
            with Image.open(annotated) as image:
                self.assertEqual(image.size, (480, 360))
            with Image.open(enhanced) as image:
                self.assertEqual(image.size, (224, 168))

    def test_enhancement_variants_are_labeled_and_comparable(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "item.jpg"
            image = Image.new("RGB", (40, 30), "white")
            for x in range(8, 32):
                image.putpixel((x, 15), (20, 20, 20))
            image.save(source)

            variants = create_enhancement_variants(source, root / "variants", scale=4)
            self.assertEqual(
                [variant["key"] for variant in variants],
                [
                    "original",
                    "lanczos_4x",
                    "clahe_color",
                    "denoise_sharpen",
                    "adaptive_threshold",
                    "otsu_threshold",
                ],
            )
            self.assertEqual([variant["use"] for variant in variants[-2:]], ["OCR only", "OCR only"])
            for variant in variants:
                self.assertTrue(variant["path"].is_file())
            with Image.open(variants[0]["path"]) as original:
                self.assertEqual(original.size, (40, 30))
            with Image.open(variants[1]["path"]) as enlarged:
                self.assertEqual(enlarged.size, (160, 120))

            comparison = create_enhancement_comparison_sheet(
                root / "comparison.png",
                [("Item 001", variants)],
            )
            with Image.open(comparison) as sheet:
                self.assertEqual(sheet.size, (1560, 300))

    def test_priced_workbook_preserves_source_and_embedded_images(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source_image = root / "source.jpg"
            Image.new("RGB", (160, 120), "navy").save(source_image)
            items = export_item_images(
                source_image,
                [ItemSelection(1, ((10, 10), (120, 90)), description="Test collectible")],
                root / "crops",
            )
            source_workbook = create_inventory_workbook(root / "inventory.xlsx", source_image, items)
            original_hash = hashlib.sha256(source_workbook.read_bytes()).hexdigest()
            observation = {
                "match_tier": "same_model_or_edition",
                "match_confidence": "high",
                "valuation_eligible": True,
                "valuation_exclusion_reasons": [],
                "sale_status": "sold",
                "price_basis": "sold_price",
                "amount": 75,
                "valuation_amount": 75,
                "currency": "USD",
                "sold_at": "2026-09-01T00:00:00Z",
                "platform": "eBay",
                "listing_or_lot_id": "123",
                "title": "Exact test collectible",
                "source_url": "https://example.test/item/123",
                "evidence_status": "supported",
                "engagement": {"watchers": 0, "likes": 4},
                "engagement_captured_at": "2026-09-20T12:00:00Z",
                "captured_at": "2026-09-20T12:00:00Z",
            }
            evaluation = {
                "observations": [observation],
                "exact_results": [observation],
                "last_sold": observation,
                "expected_resale": {
                    "low": 75,
                    "median": 75,
                    "high": 75,
                    "sold_count": 1,
                    "match_tier": "same_model_or_edition",
                    "confidence": "low",
                },
                "evidence_summary": {"exact": 1, "valuation_eligible_sold": 1},
            }
            output = apply_price_evaluations_to_workbook(
                source_workbook,
                root / "inventory_priced.xlsx",
                {1: evaluation},
                evidence_json_path=root / "inventory_priced.json",
            )
            self.assertEqual(hashlib.sha256(source_workbook.read_bytes()).hexdigest(), original_hash)
            workbook = load_workbook(output, read_only=False, data_only=False)
            inventory = workbook["Inventory"]
            headings = {cell.value: cell.column for cell in inventory[1]}
            self.assertEqual(inventory.cell(2, headings["Pricing Status"]).value, "Exact sold evidence")
            self.assertEqual(inventory.cell(2, headings["Expected Resale Median"]).value, 75)
            self.assertIn("0 watchers", inventory.cell(2, headings["Matching Engagement"]).value)
            evidence = workbook["Price Evidence"]
            evidence_headings = {cell.value: cell.column for cell in evidence[1]}
            self.assertEqual(evidence.cell(2, evidence_headings["Watchers"]).value, 0)
            self.assertIsNone(evidence.cell(2, evidence_headings["Bids"]).value)
            self.assertEqual(evidence.cell(2, evidence_headings["Source URL"]).hyperlink.target, "https://example.test/item/123")
            self.assertEqual(len(inventory._images), 1)
            workbook.close()
            self.assertTrue((root / "inventory_priced.json").is_file())


if __name__ == "__main__":
    unittest.main()
