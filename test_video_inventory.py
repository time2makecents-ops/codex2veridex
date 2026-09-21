import hashlib
import json
import tempfile
import unittest
from pathlib import Path

import cv2
import numpy as np
from openpyxl import load_workbook
from PIL import Image

from video_inventory import (
    VideoInventorySelection,
    create_video_item_proposals,
    create_workbook_from_video_candidates,
    extract_video_inventory_candidates,
    propose_item_lassos,
)


class VideoInventoryTests(unittest.TestCase):
    @staticmethod
    def _make_video(path: Path) -> None:
        writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"MJPG"), 5.0, (320, 240))
        if not writer.isOpened():
            raise unittest.SkipTest("MJPG test-video writer is unavailable")
        try:
            for index in range(30):
                frame = np.full((240, 320, 3), 238, dtype=np.uint8)
                scene = index // 10
                colors = ((30, 70, 210), (60, 175, 60), (190, 70, 50))
                left = 35 + scene * 18
                cv2.rectangle(frame, (left, 45), (left + 170, 205), colors[scene], -1)
                cv2.circle(frame, (left + 85, 125), 38 + scene * 4, (245, 245, 245), 5)
                cv2.putText(frame, f"ITEM {scene + 1}", (left + 28, 132), cv2.FONT_HERSHEY_SIMPLEX, 0.65, (15, 15, 15), 2)
                writer.write(frame)
        finally:
            writer.release()

    def test_extracts_review_candidates_then_builds_selected_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            video = root / "items.avi"
            self._make_video(video)
            source_hash = hashlib.sha256(video.read_bytes()).hexdigest()

            result = extract_video_inventory_candidates(
                video,
                root / "video_project",
                sample_seconds=0.4,
                window_seconds=2.0,
                max_candidates=10,
            )
            self.assertEqual(result["schema"], "veridex.video_inventory.v1")
            self.assertTrue(result["review_required"])
            self.assertGreaterEqual(len(result["candidates"]), 2)
            self.assertEqual(hashlib.sha256(video.read_bytes()).hexdigest(), source_hash)
            manifest_path = Path(result["manifest_path"])
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            contact_sheet = manifest_path.parent / manifest["contact_sheet"]
            with Image.open(contact_sheet) as sheet:
                self.assertGreater(sheet.width, 0)
                self.assertGreater(sheet.height, 0)
            for candidate in manifest["candidates"]:
                image_path = manifest_path.parent / candidate["image_path"]
                self.assertEqual(hashlib.sha256(image_path.read_bytes()).hexdigest(), candidate["image_sha256"])

            first = manifest["candidates"][0]
            output = create_workbook_from_video_candidates(
                manifest_path,
                [VideoInventorySelection(
                    first["candidate_id"],
                    points=((25, 35), (245, 220)),
                    description="Approved video item",
                )],
                root / "video_inventory.xlsx",
                root / "approved_crops",
            )
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                self.assertEqual(workbook["Inventory"]["C2"].value, "Approved video item")
                self.assertIn(first["candidate_id"], workbook["Inventory"]["R2"].value)
                self.assertEqual(len(workbook["Inventory"]._images), 1)
            finally:
                workbook.close()

    def test_requires_new_project_and_honors_cancellation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            video = root / "items.avi"
            self._make_video(video)
            with self.assertRaisesRegex(RuntimeError, "canceled"):
                extract_video_inventory_candidates(video, root / "cancelled", cancel_check=lambda: True)
            result = extract_video_inventory_candidates(video, root / "project")
            self.assertTrue(Path(result["manifest_path"]).is_file())
            with self.assertRaises(FileExistsError):
                extract_video_inventory_candidates(video, root / "project")

    def test_item_lasso_proposals_cover_whole_regions_and_flag_frame_edges(self) -> None:
        image = Image.new("RGB", (640, 420), "white")
        array = np.asarray(image).copy()
        cv2.rectangle(array, (75, 85), (245, 330), (220, 65, 45), -1)
        cv2.rectangle(array, (75, 85), (245, 330), (15, 15, 15), 6)
        cv2.circle(array, (440, 210), 95, (40, 155, 70), -1)
        cv2.circle(array, (440, 210), 95, (15, 15, 15), 6)
        proposals = propose_item_lassos(Image.fromarray(array), maximum_proposals=6)
        self.assertGreaterEqual(len(proposals), 2)
        bounds = [proposal["bounds"] for proposal in proposals]
        self.assertTrue(any(left <= 80 and top <= 90 and right >= 240 and bottom >= 325 for left, top, right, bottom in bounds))
        self.assertTrue(any(left <= 350 and top <= 120 and right >= 530 and bottom >= 300 for left, top, right, bottom in bounds))

        edge_array = np.full((320, 480, 3), 245, dtype=np.uint8)
        cv2.rectangle(edge_array, (0, 55), (190, 270), (40, 80, 210), -1)
        cv2.rectangle(edge_array, (0, 55), (190, 270), (15, 15, 15), 7)
        edge_proposals = propose_item_lassos(Image.fromarray(edge_array), maximum_proposals=4)
        self.assertTrue(any("touches_frame_edge" in proposal["quality_flags"] for proposal in edge_proposals))

    def test_review_proposals_preserve_provenance_and_require_workbook_review(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            video = root / "items.avi"
            self._make_video(video)
            extracted = extract_video_inventory_candidates(
                video,
                root / "project",
                sample_seconds=0.4,
                window_seconds=2.0,
                max_candidates=10,
            )
            first_id = extracted["candidates"][0]["candidate_id"]
            result = create_video_item_proposals(
                Path(extracted["manifest_path"]),
                root / "proposals",
                rotations={first_id: 90},
                maximum_proposals_per_scene=4,
            )
            self.assertEqual(result["schema"], "veridex.video_inventory.item_proposals.v1")
            self.assertFalse(result["workbook_ready"])
            self.assertTrue(result["review_required"])
            self.assertEqual(result["source_video_sha256"], extracted["source"]["sha256"])
            self.assertTrue(Path(result["manifest_path"]).is_file())
            orientation_review = Path(result["output_dir"]) / result["orientation_review"]
            proposal_sheet = Path(result["output_dir"]) / result["proposal_contact_sheet"]
            self.assertTrue(orientation_review.is_file())
            self.assertTrue(proposal_sheet.is_file())
            self.assertEqual(hashlib.sha256(orientation_review.read_bytes()).hexdigest(), result["orientation_review_sha256"])
            self.assertEqual(hashlib.sha256(proposal_sheet.read_bytes()).hexdigest(), result["proposal_contact_sheet_sha256"])
            rotated = next(
                scene for scene in result["scenes"]
                if first_id in scene["candidate_ids"]
            )
            self.assertEqual(rotated["rotation_degrees"], 90)
            self.assertEqual(rotated["orientation_status"], "accepted_override")
            with self.assertRaises(FileExistsError):
                create_video_item_proposals(Path(extracted["manifest_path"]), root / "proposals")

    def test_multiple_rotated_regions_may_be_approved_from_one_view(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            video = root / "items.avi"
            self._make_video(video)
            extracted = extract_video_inventory_candidates(video, root / "project")
            candidate_id = extracted["candidates"][0]["candidate_id"]
            output = create_workbook_from_video_candidates(
                Path(extracted["manifest_path"]),
                [
                    VideoInventorySelection(candidate_id, points=((20, 20), (130, 180)), description="Region one", rotation_degrees=90),
                    VideoInventorySelection(candidate_id, points=((80, 40), (220, 200)), description="Region two", rotation_degrees=90),
                ],
                root / "two_regions.xlsx",
                root / "two_regions_crops",
            )
            workbook = load_workbook(output, read_only=False, data_only=False)
            try:
                self.assertEqual(workbook["Inventory"]["C2"].value, "Region one")
                self.assertEqual(workbook["Inventory"]["C3"].value, "Region two")
                self.assertIn("Rotation 90 degrees clockwise", workbook["Inventory"]["R2"].value)
                self.assertEqual(len(workbook["Inventory"]._images), 2)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
