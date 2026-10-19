"""Private batch-and-item intake catalog for the paired Veridex browser app.

The SQLite catalog is the searchable authority.  Each item also has a generated,
human-readable ``item.json`` snapshot beside its source files so an item can be
inspected or moved as one self-contained folder without making file paths the
application database.
"""

from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import shutil
import sqlite3
import threading
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterator, Optional
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ITEM_STATES = {"intake", "needs_review", "research_ready", "handed_off", "listed", "archived"}
EVIDENCE_STATES = {"note", "accepted", "needs_review"}
ASSET_ROLES = {"attachment", "photo", "research_source", "document", "media"}
SAFE_ID_RE = re.compile(r"[A-Za-z0-9_-]+")


class SellerIntakeError(RuntimeError):
    """A bounded, safe-to-display Item Intake error."""


def _now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _identifier(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:16]}"


def _text(value: Any, limit: int) -> str:
    return " ".join(str(value or "").split())[:limit]


def _body(value: Any, limit: int) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()[:limit]


def _slug(value: Any, fallback: str = "item") -> str:
    clean = re.sub(r"[^a-z0-9]+", "-", _text(value, 100).lower()).strip("-")
    return clean[:64] or fallback


class SellerIntakeCatalog:
    """Workspace-independent local intake catalog with contained file storage."""

    def __init__(self, data_root: Path) -> None:
        self.root = Path(data_root).resolve() / "seller_intake"
        self.db_path = self.root / "catalog.sqlite3"
        self.batches_root = self.root / "batches"
        self._lock = threading.RLock()
        self.root.mkdir(parents=True, exist_ok=True)
        self.batches_root.mkdir(parents=True, exist_ok=True)
        self._ensure_schema()

    @contextmanager
    def _connection(self) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(str(self.db_path), timeout=20, check_same_thread=False)
        connection.row_factory = sqlite3.Row
        try:
            yield connection
            connection.commit()
        finally:
            connection.close()

    def _ensure_schema(self) -> None:
        with self._connection() as connection:
            connection.executescript(
                """
                PRAGMA foreign_keys = ON;
                CREATE TABLE IF NOT EXISTS batches (
                    batch_id TEXT PRIMARY KEY,
                    title TEXT NOT NULL,
                    notes TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS items (
                    item_id TEXT PRIMARY KEY,
                    batch_id TEXT NOT NULL REFERENCES batches(batch_id) ON DELETE CASCADE,
                    title TEXT NOT NULL DEFAULT '',
                    description TEXT NOT NULL DEFAULT '',
                    estimated_value REAL,
                    currency TEXT NOT NULL DEFAULT 'USD',
                    research_needed INTEGER NOT NULL DEFAULT 0,
                    status TEXT NOT NULL DEFAULT 'intake',
                    notes TEXT NOT NULL DEFAULT '',
                    reference_id TEXT NOT NULL DEFAULT '',
                    chatgpt_thread_url TEXT NOT NULL DEFAULT '',
                    gemini_thread_url TEXT NOT NULL DEFAULT '',
                    tags_json TEXT NOT NULL DEFAULT '[]',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS assets (
                    asset_id TEXT PRIMARY KEY,
                    item_id TEXT NOT NULL REFERENCES items(item_id) ON DELETE CASCADE,
                    source_asset_id TEXT REFERENCES assets(asset_id),
                    transformation_json TEXT NOT NULL DEFAULT '{}',
                    role TEXT NOT NULL,
                    name TEXT NOT NULL,
                    original_name TEXT NOT NULL,
                    content_type TEXT NOT NULL,
                    size INTEGER NOT NULL,
                    sha256 TEXT NOT NULL,
                    relative_path TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS evidence (
                    evidence_id TEXT PRIMARY KEY,
                    item_id TEXT NOT NULL REFERENCES items(item_id) ON DELETE CASCADE,
                    kind TEXT NOT NULL,
                    claim TEXT NOT NULL,
                    source_asset_id TEXT,
                    source_excerpt TEXT NOT NULL DEFAULT '',
                    observed_date TEXT NOT NULL DEFAULT '',
                    location_text TEXT NOT NULL DEFAULT '',
                    confidence TEXT NOT NULL DEFAULT 'unrated',
                    status TEXT NOT NULL DEFAULT 'note',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS handoffs (
                    handoff_id TEXT PRIMARY KEY,
                    item_id TEXT NOT NULL REFERENCES items(item_id) ON DELETE CASCADE,
                    workspace_id TEXT NOT NULL,
                    session_id TEXT NOT NULL,
                    artifact_file_ids_json TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS item_batch_idx ON items(batch_id, updated_at DESC);
                CREATE INDEX IF NOT EXISTS asset_item_idx ON assets(item_id, created_at DESC);
                CREATE INDEX IF NOT EXISTS evidence_item_idx ON evidence(item_id, created_at DESC);
                """
            )
            item_columns = {str(row["name"]) for row in connection.execute("PRAGMA table_info(items)").fetchall()}
            if "reference_id" not in item_columns:
                connection.execute("ALTER TABLE items ADD COLUMN reference_id TEXT NOT NULL DEFAULT ''")
            if "chatgpt_thread_url" not in item_columns:
                connection.execute("ALTER TABLE items ADD COLUMN chatgpt_thread_url TEXT NOT NULL DEFAULT ''")
            if "gemini_thread_url" not in item_columns:
                connection.execute("ALTER TABLE items ADD COLUMN gemini_thread_url TEXT NOT NULL DEFAULT ''")
            asset_columns = {str(row["name"]) for row in connection.execute("PRAGMA table_info(assets)").fetchall()}
            if "source_asset_id" not in asset_columns:
                connection.execute("ALTER TABLE assets ADD COLUMN source_asset_id TEXT REFERENCES assets(asset_id)")
            if "transformation_json" not in asset_columns:
                connection.execute("ALTER TABLE assets ADD COLUMN transformation_json TEXT NOT NULL DEFAULT '{}'")

    @staticmethod
    def _safe_id(value: Any, label: str) -> str:
        candidate = str(value or "").strip()
        if not SAFE_ID_RE.fullmatch(candidate):
            raise SellerIntakeError(f"{label} is invalid")
        return candidate

    @staticmethod
    def _safe_filename(value: Any) -> str:
        name = Path(str(value or "")).name.strip()
        name = re.sub(r"[<>:\\|?*\x00-\x1f]", "_", name).strip(". ")
        if not name:
            raise SellerIntakeError("A file name is required")
        return name[:180]

    def _batch_row(self, connection: sqlite3.Connection, batch_id: str) -> sqlite3.Row:
        row = connection.execute("SELECT * FROM batches WHERE batch_id = ?", (self._safe_id(batch_id, "batch_id"),)).fetchone()
        if not row:
            raise KeyError("Unknown intake batch")
        return row

    def _item_row(self, connection: sqlite3.Connection, item_id: str) -> sqlite3.Row:
        row = connection.execute("SELECT * FROM items WHERE item_id = ?", (self._safe_id(item_id, "item_id"),)).fetchone()
        if not row:
            raise KeyError("Unknown intake item")
        return row

    def _item_dir(self, batch_id: str, item_id: str, title: str) -> Path:
        return self.batches_root / batch_id / "items" / f"{item_id}_{_slug(title, 'untitled-item')}"

    def _find_item_dir(self, connection: sqlite3.Connection, item: sqlite3.Row) -> Path:
        base = self.batches_root / str(item["batch_id"]) / "items"
        prefix = f"{item['item_id']}_"
        matches = [path for path in base.glob(f"{prefix}*") if path.is_dir()] if base.is_dir() else []
        if matches:
            return matches[0]
        return self._item_dir(str(item["batch_id"]), str(item["item_id"]), str(item["title"]))

    def _public_batch(self, row: sqlite3.Row, count: int = 0) -> Dict[str, Any]:
        return {**dict(row), "item_count": int(count)}

    @staticmethod
    def _public_asset(row: sqlite3.Row) -> Dict[str, Any]:
        result = {key: row[key] for key in (
            "asset_id", "item_id", "role", "name", "original_name", "content_type", "size", "sha256", "created_at"
        )}
        result["source_asset_id"] = row["source_asset_id"] if "source_asset_id" in row.keys() else None
        try:
            result["transformation"] = json.loads(str(row["transformation_json"] or "{}")) if "transformation_json" in row.keys() else {}
        except json.JSONDecodeError:
            result["transformation"] = {}
        return result

    @staticmethod
    def _public_evidence(row: sqlite3.Row) -> Dict[str, Any]:
        return dict(row)

    def _public_item(self, connection: sqlite3.Connection, item: sqlite3.Row, *, detail: bool = True) -> Dict[str, Any]:
        result = dict(item)
        try:
            result["tags"] = json.loads(str(result.pop("tags_json") or "[]"))
        except json.JSONDecodeError:
            result["tags"] = []
        result["asset_count"] = int(connection.execute("SELECT COUNT(*) FROM assets WHERE item_id = ?", (item["item_id"],)).fetchone()[0])
        result["evidence_count"] = int(connection.execute("SELECT COUNT(*) FROM evidence WHERE item_id = ?", (item["item_id"],)).fetchone()[0])
        if detail:
            assets = connection.execute("SELECT * FROM assets WHERE item_id = ? ORDER BY created_at", (item["item_id"],)).fetchall()
            evidence = connection.execute("SELECT * FROM evidence WHERE item_id = ? ORDER BY created_at DESC", (item["item_id"],)).fetchall()
            result["assets"] = [self._public_asset(row) for row in assets]
            result["evidence"] = [self._public_evidence(row) for row in evidence]
        return result

    def _write_snapshot(self, connection: sqlite3.Connection, item_id: str) -> None:
        item = self._item_row(connection, item_id)
        item_dir = self._find_item_dir(connection, item)
        item_dir.mkdir(parents=True, exist_ok=True)
        for child in ("originals", "derivatives", "research_sources", "exports"):
            (item_dir / child).mkdir(exist_ok=True)
        snapshot = self._public_item(connection, item, detail=True)
        snapshot["folder_layout"] = {"originals": "originals", "derivatives": "derivatives", "research_sources": "research_sources", "exports": "exports"}
        snapshot_path = item_dir / "item.json"
        temporary = snapshot_path.with_suffix(".json.tmp")
        temporary.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(snapshot_path)

    def create_batch(self, value: Dict[str, Any]) -> Dict[str, Any]:
        title = _text(value.get("title"), 160) or f"Intake {datetime.now().strftime('%Y-%m-%d')}"
        notes = _body(value.get("notes"), 4000)
        now = _now()
        batch_id = _identifier("batch")
        with self._lock, self._connection() as connection:
            connection.execute("INSERT INTO batches(batch_id, title, notes, created_at, updated_at) VALUES(?, ?, ?, ?, ?)", (batch_id, title, notes, now, now))
        (self.batches_root / batch_id / "items").mkdir(parents=True, exist_ok=True)
        return self.get_batch(batch_id)

    def list_batches(self, *, include_archived: bool = True) -> list[Dict[str, Any]]:
        item_count = "COUNT(items.item_id)" if include_archived else "COUNT(CASE WHEN items.status != 'archived' THEN items.item_id END)"
        with self._connection() as connection:
            rows = connection.execute(
                f"SELECT batches.*, {item_count} AS item_count FROM batches LEFT JOIN items ON items.batch_id = batches.batch_id GROUP BY batches.batch_id ORDER BY batches.updated_at DESC"
            ).fetchall()
        return [self._public_batch(row, int(row["item_count"])) for row in rows]

    def get_batch(self, batch_id: str, *, include_archived: bool = True) -> Dict[str, Any]:
        with self._connection() as connection:
            batch = self._batch_row(connection, batch_id)
            query = "SELECT * FROM items WHERE batch_id = ?"
            if not include_archived:
                query += " AND status != 'archived'"
            items = connection.execute(query + " ORDER BY updated_at DESC", (batch["batch_id"],)).fetchall()
            return {**self._public_batch(batch, len(items)), "items": [self._public_item(connection, row, detail=False) for row in items]}

    def create_item(self, value: Dict[str, Any]) -> Dict[str, Any]:
        batch_id = self._safe_id(value.get("batch_id"), "batch_id")
        title = _text(value.get("title"), 200)
        description = _body(value.get("description"), 10000)
        currency = _text(value.get("currency"), 12).upper() or "USD"
        status = str(value.get("status") or "intake").strip().lower()
        if status not in ITEM_STATES:
            raise SellerIntakeError("Item status is invalid")
        estimated_value = self._value(value.get("estimated_value"))
        tags = self._tags(value.get("tags"))
        reference_id = _text(value.get("reference_id"), 120)
        chatgpt_thread_url = self._chatgpt_thread_url(value.get("chatgpt_thread_url"))
        gemini_thread_url = self._gemini_thread_url(value.get("gemini_thread_url"))
        now = _now()
        item_id = _identifier("item")
        with self._lock, self._connection() as connection:
            self._batch_row(connection, batch_id)
            connection.execute(
                "INSERT INTO items(item_id, batch_id, title, description, estimated_value, currency, research_needed, status, notes, reference_id, chatgpt_thread_url, gemini_thread_url, tags_json, created_at, updated_at) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (item_id, batch_id, title, description, estimated_value, currency, int(bool(value.get("research_needed"))), status, _body(value.get("notes"), 10000), reference_id, chatgpt_thread_url, gemini_thread_url, json.dumps(tags), now, now),
            )
            connection.execute("UPDATE batches SET updated_at = ? WHERE batch_id = ?", (now, batch_id))
            self._write_snapshot(connection, item_id)
            row = self._item_row(connection, item_id)
            return self._public_item(connection, row)

    @staticmethod
    def _value(value: Any) -> Optional[float]:
        if value in (None, ""):
            return None
        try:
            parsed = round(float(value), 2)
        except (TypeError, ValueError) as exc:
            raise SellerIntakeError("Estimated value must be a number") from exc
        if parsed < 0 or parsed >= 100_000_000:
            raise SellerIntakeError("Estimated value is outside the supported range")
        return parsed

    @staticmethod
    def _tags(value: Any) -> list[str]:
        raw = value.split(",") if isinstance(value, str) else (value if isinstance(value, list) else [])
        return list(dict.fromkeys(_text(row, 80) for row in raw if _text(row, 80)))[:20]

    @staticmethod
    def _chatgpt_thread_url(value: Any) -> str:
        url = _text(value, 2000)
        if not url:
            return ""
        parsed = urlparse(url)
        if parsed.scheme != "https" or parsed.netloc.lower() != "chatgpt.com" or not parsed.path.startswith("/share/"):
            raise SellerIntakeError("ChatGPT thread link must be an https://chatgpt.com/share/... URL")
        return url

    @staticmethod
    def _gemini_thread_url(value: Any) -> str:
        url = _text(value, 2000)
        if not url:
            return ""
        parsed = urlparse(url)
        if parsed.scheme != "https" or parsed.netloc.lower() != "g.co" or not parsed.path.startswith("/gemini/share/"):
            raise SellerIntakeError("Gemini thread link must be an https://g.co/gemini/share/... URL")
        return url

    def get_item(self, item_id: str) -> Dict[str, Any]:
        with self._connection() as connection:
            return self._public_item(connection, self._item_row(connection, item_id))

    def update_item(self, value: Dict[str, Any]) -> Dict[str, Any]:
        item_id = self._safe_id(value.get("item_id"), "item_id")
        allowed = {"title", "description", "estimated_value", "currency", "research_needed", "status", "notes", "reference_id", "chatgpt_thread_url", "gemini_thread_url", "tags"}
        updates: Dict[str, Any] = {}
        for key in allowed:
            if key not in value:
                continue
            if key == "title": updates[key] = _text(value[key], 200)
            elif key == "description" or key == "notes": updates[key] = _body(value[key], 10000)
            elif key == "reference_id": updates[key] = _text(value[key], 120)
            elif key == "estimated_value": updates[key] = self._value(value[key])
            elif key == "currency": updates[key] = _text(value[key], 12).upper() or "USD"
            elif key == "chatgpt_thread_url": updates[key] = self._chatgpt_thread_url(value[key])
            elif key == "gemini_thread_url": updates[key] = self._gemini_thread_url(value[key])
            elif key == "research_needed": updates[key] = int(bool(value[key]))
            elif key == "status":
                status = str(value[key] or "").strip().lower()
                if status not in ITEM_STATES: raise SellerIntakeError("Item status is invalid")
                updates[key] = status
            elif key == "tags": updates["tags_json"] = json.dumps(self._tags(value[key]))
        if not updates:
            raise SellerIntakeError("No item fields were supplied")
        now = _now()
        with self._lock, self._connection() as connection:
            item = self._item_row(connection, item_id)
            columns = list(updates) + ["updated_at"]
            values = [updates[key] for key in updates] + [now, item_id]
            connection.execute(f"UPDATE items SET {', '.join(f'{column} = ?' for column in columns)} WHERE item_id = ?", values)
            connection.execute("UPDATE batches SET updated_at = ? WHERE batch_id = ?", (now, item["batch_id"]))
            self._write_snapshot(connection, item_id)
            return self._public_item(connection, self._item_row(connection, item_id))

    def delete_item(self, item_id: str) -> Dict[str, str]:
        """Permanently remove one intake item and its private source folder."""
        item_id = self._safe_id(item_id, "item_id")
        with self._lock, self._connection() as connection:
            item = self._item_row(connection, item_id)
            item_dir = self._find_item_dir(connection, item).resolve()
            items_root = (self.batches_root / str(item["batch_id"]) / "items").resolve()
            if items_root not in item_dir.parents:
                raise SellerIntakeError("Item folder is outside the intake storage")
            if item_dir.exists():
                shutil.rmtree(item_dir)
            connection.execute("DELETE FROM handoffs WHERE item_id = ?", (item_id,))
            connection.execute("DELETE FROM evidence WHERE item_id = ?", (item_id,))
            connection.execute("DELETE FROM assets WHERE item_id = ?", (item_id,))
            connection.execute("DELETE FROM items WHERE item_id = ?", (item_id,))
            connection.execute("UPDATE batches SET updated_at = ? WHERE batch_id = ?", (_now(), item["batch_id"]))
            return {"item_id": item_id, "batch_id": str(item["batch_id"]), "title": str(item["title"])}

    def add_asset(
        self,
        item_id: str,
        filename: str,
        content: bytes,
        content_type: str = "",
        role: str = "attachment",
        *,
        source_asset_id: str = "",
        transformation: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        if not content:
            raise SellerIntakeError("File is empty")
        role = str(role or "attachment").strip().lower()
        if role not in ASSET_ROLES:
            raise SellerIntakeError("Asset role is invalid")
        original_name = self._safe_filename(filename)
        asset_id = _identifier("asset")
        now = _now()
        with self._lock, self._connection() as connection:
            item = self._item_row(connection, item_id)
            source_asset_id = str(source_asset_id or "").strip()
            if source_asset_id:
                source = self._asset_row(connection, source_asset_id)
                if str(source["item_id"]) != str(item["item_id"]):
                    raise SellerIntakeError("Derived photo source does not belong to this item")
            item_dir = self._find_item_dir(connection, item)
            bucket = "research_sources" if role == "research_source" else ("derivatives" if source_asset_id else "originals")
            folder = item_dir / bucket
            folder.mkdir(parents=True, exist_ok=True)
            stored_name = f"{asset_id}__{original_name}"
            path = folder / stored_name
            path.write_bytes(content)
            relative_path = path.relative_to(self.root).as_posix()
            sha256 = hashlib.sha256(content).hexdigest()
            content_type = str(content_type or mimetypes.guess_type(original_name)[0] or "application/octet-stream")[:200]
            connection.execute(
                "INSERT INTO assets(asset_id, item_id, source_asset_id, transformation_json, role, name, original_name, content_type, size, sha256, relative_path, created_at) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (asset_id, item_id, source_asset_id or None, json.dumps(transformation or {}, ensure_ascii=False, sort_keys=True), role, original_name, original_name, content_type, len(content), sha256, relative_path, now),
            )
            connection.execute("UPDATE items SET updated_at = ? WHERE item_id = ?", (now, item_id))
            connection.execute("UPDATE batches SET updated_at = ? WHERE batch_id = ?", (now, item["batch_id"]))
            self._write_snapshot(connection, item_id)
            row = connection.execute("SELECT * FROM assets WHERE asset_id = ?", (asset_id,)).fetchone()
            return self._public_asset(row)

    def delete_asset(self, item_id: str, asset_id: str) -> Dict[str, str]:
        """Permanently remove one asset owned by an intake item."""
        item_id = self._safe_id(item_id, "item_id")
        asset_id = self._safe_id(asset_id, "asset_id")
        with self._lock, self._connection() as connection:
            item = self._item_row(connection, item_id)
            asset = self._asset_row(connection, asset_id)
            if str(asset["item_id"]) != item_id:
                raise SellerIntakeError("Asset does not belong to this item")
            if connection.execute("SELECT 1 FROM evidence WHERE source_asset_id = ? LIMIT 1", (asset_id,)).fetchone():
                raise SellerIntakeError("This file is used by a finding and cannot be deleted")
            if connection.execute("SELECT 1 FROM assets WHERE source_asset_id = ? LIMIT 1", (asset_id,)).fetchone():
                raise SellerIntakeError("Remove derived photos before deleting their source photo")
            path = (self.root / str(asset["relative_path"])).resolve()
            if self.root not in path.parents or not path.is_file():
                raise SellerIntakeError("Stored asset is unavailable")
            path.unlink()
            connection.execute("DELETE FROM assets WHERE asset_id = ?", (asset_id,))
            connection.execute("UPDATE items SET updated_at = ? WHERE item_id = ?", (_now(), item_id))
            connection.execute("UPDATE batches SET updated_at = ? WHERE batch_id = ?", (_now(), item["batch_id"]))
            self._write_snapshot(connection, item_id)
            return {"asset_id": asset_id, "item_id": item_id, "name": str(asset["name"])}

    def add_source_text(self, value: Dict[str, Any]) -> Dict[str, Any]:
        item_id = self._safe_id(value.get("item_id"), "item_id")
        text = _body(value.get("text"), 1_000_000)
        if not text:
            raise SellerIntakeError("Research text is required")
        label = _text(value.get("title"), 160) or "research-chat"
        return self.add_asset(item_id, f"{_slug(label, 'research-chat')}.txt", text.encode("utf-8"), "text/plain", "research_source")

    def add_evidence(self, value: Dict[str, Any]) -> Dict[str, Any]:
        item_id = self._safe_id(value.get("item_id"), "item_id")
        claim = _body(value.get("claim"), 5000)
        if not claim:
            raise SellerIntakeError("Evidence claim is required")
        status = str(value.get("status") or "note").strip().lower()
        if status not in EVIDENCE_STATES:
            raise SellerIntakeError("Evidence status is invalid")
        source_asset_id = str(value.get("source_asset_id") or "").strip()
        now = _now()
        evidence_id = _identifier("evidence")
        with self._lock, self._connection() as connection:
            self._item_row(connection, item_id)
            if source_asset_id:
                asset = connection.execute("SELECT asset_id FROM assets WHERE asset_id = ? AND item_id = ?", (source_asset_id, item_id)).fetchone()
                if not asset: raise SellerIntakeError("Evidence source does not belong to this item")
            connection.execute(
                "INSERT INTO evidence(evidence_id, item_id, kind, claim, source_asset_id, source_excerpt, observed_date, location_text, confidence, status, created_at, updated_at) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (evidence_id, item_id, _text(value.get("kind"), 80) or "note", claim, source_asset_id or None, _body(value.get("source_excerpt"), 4000), _text(value.get("observed_date"), 80), _text(value.get("location_text"), 240), _text(value.get("confidence"), 80) or "unrated", status, now, now),
            )
            connection.execute("UPDATE items SET updated_at = ? WHERE item_id = ?", (now, item_id))
            self._write_snapshot(connection, item_id)
            return self._public_evidence(connection.execute("SELECT * FROM evidence WHERE evidence_id = ?", (evidence_id,)).fetchone())

    def _asset_row(self, connection: sqlite3.Connection, asset_id: str) -> sqlite3.Row:
        row = connection.execute("SELECT * FROM assets WHERE asset_id = ?", (self._safe_id(asset_id, "asset_id"),)).fetchone()
        if not row: raise KeyError("Unknown intake asset")
        return row

    def asset_path(self, asset_id: str) -> tuple[sqlite3.Row, Path]:
        with self._connection() as connection:
            row = self._asset_row(connection, asset_id)
            path = (self.root / str(row["relative_path"])).resolve()
            if self.root not in path.parents or not path.is_file():
                raise KeyError("Intake asset is unavailable")
            return row, path

    def export_path(self, relative_path: str) -> Path:
        candidate = (self.root / str(relative_path or "")).resolve()
        exports_root = (self.batches_root).resolve()
        if exports_root not in candidate.parents or not candidate.is_file() or candidate.suffix.lower() not in {".csv", ".xlsx"}:
            raise KeyError("Unknown intake export")
        return candidate

    def export_batch(self, batch_id: str) -> Dict[str, Any]:
        with self._lock, self._connection() as connection:
            batch = self._batch_row(connection, batch_id)
            rows = connection.execute("SELECT * FROM items WHERE batch_id = ? ORDER BY created_at", (batch_id,)).fetchall()
            export_dir = self.batches_root / batch_id / "exports"
            export_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            headers = ["Item ID", "Reference ID", "Title", "Description", "Estimated Value", "Currency", "Research Needed", "Status", "Tags", "Notes", "ChatGPT Thread Link", "Gemini Thread Link", "Assets", "Evidence"]
            values = []
            for item in rows:
                public = self._public_item(connection, item)
                values.append([item["item_id"], item["reference_id"], item["title"], item["description"], item["estimated_value"], item["currency"], bool(item["research_needed"]), item["status"], ", ".join(public["tags"]), item["notes"], item["chatgpt_thread_url"], item["gemini_thread_url"], public["asset_count"], public["evidence_count"]])
            csv_path = export_dir / f"inventory_{stamp}.csv"
            with csv_path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream); writer.writerow(headers); writer.writerows(values)
            xlsx_path = export_dir / f"inventory_{stamp}.xlsx"
            workbook = Workbook(); sheet = workbook.active; sheet.title = "Inventory"; sheet.append(headers)
            for cell in sheet[1]: cell.font = Font(color="FFFFFF", bold=True); cell.fill = PatternFill("solid", fgColor="1F4E78")
            for row in values: sheet.append(row)
            sheet.freeze_panes = "A2"; sheet.auto_filter.ref = f"A1:N{max(1, sheet.max_row)}"
            for column, width in enumerate((22, 22, 28, 46, 16, 11, 16, 18, 24, 40, 48, 48, 10, 10), start=1): sheet.column_dimensions[chr(64 + column)].width = width
            workbook.save(xlsx_path)
            return {"batch": self._public_batch(batch, len(rows)), "exports": [{"name": csv_path.name, "relative_path": csv_path.relative_to(self.root).as_posix(), "content_type": "text/csv", "size": csv_path.stat().st_size}, {"name": xlsx_path.name, "relative_path": xlsx_path.relative_to(self.root).as_posix(), "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "size": xlsx_path.stat().st_size}]}

    def handoff_to_veridex(self, item_id: str, workspace_id: str, session_id: str, store: Any) -> Dict[str, Any]:
        with self._lock, self._connection() as connection:
            item = self._item_row(connection, item_id)
            public = self._public_item(connection, item)
            saved = []
            for asset in public["assets"]:
                source_row = self._asset_row(connection, asset["asset_id"])
                source_path = (self.root / str(source_row["relative_path"])).resolve()
                if not source_path.is_file(): raise SellerIntakeError(f"Stored asset is unavailable: {asset['name']}")
                stored = store.save_file(workspace_id, session_id, str(asset["name"]), source_path.read_bytes(), str(asset["content_type"]), source="seller_intake_handoff", kind="upload", description=f"Seller Intake item {item_id}", metadata={"seller_intake": {"item_id": item_id, "batch_id": item["batch_id"], "asset_id": asset["asset_id"], "source_sha256": asset["sha256"]}})
                saved.append(stored)
            brief = {
                "seller_intake_item": {key: public.get(key) for key in ("item_id", "batch_id", "reference_id", "title", "description", "estimated_value", "currency", "research_needed", "status", "notes", "chatgpt_thread_url", "gemini_thread_url", "tags")},
                "evidence": public["evidence"],
                "source_assets": [{key: asset.get(key) for key in ("asset_id", "name", "content_type", "sha256")} for asset in public["assets"]],
                "instruction": "These are source materials from Seller Intake. Treat timeline entries as source-linked notes, not authentication or valuation facts. External research still requires the normal Veridex confirmation path.",
            }
            brief_file = store.save_file(workspace_id, session_id, f"{item_id}_research_brief.json", json.dumps(brief, ensure_ascii=False, indent=2).encode("utf-8"), "application/json", source="seller_intake_handoff", kind="upload", description=f"Seller Intake research brief for {item_id}", metadata={"seller_intake": {"item_id": item_id, "batch_id": item["batch_id"]}})
            saved.append(brief_file)
            now = _now(); handoff_id = _identifier("handoff")
            connection.execute("INSERT INTO handoffs(handoff_id, item_id, workspace_id, session_id, artifact_file_ids_json, created_at) VALUES(?, ?, ?, ?, ?, ?)", (handoff_id, item_id, workspace_id, session_id, json.dumps([row["file_id"] for row in saved]), now))
            connection.execute("UPDATE items SET status = 'handed_off', updated_at = ? WHERE item_id = ?", (now, item_id))
            self._write_snapshot(connection, item_id)
            return {"handoff_id": handoff_id, "item_id": item_id, "workspace_id": workspace_id, "session_id": session_id, "files": [{key: row.get(key) for key in ("file_id", "artifact_number", "name", "content_type", "size", "sha256")} for row in saved]}
