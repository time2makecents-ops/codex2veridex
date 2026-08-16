"""Governed Resume Studio storage, analysis, validation, and document rendering."""

from __future__ import annotations

import io
import html
import ipaddress
import json
import re
import socket
import urllib.request
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List


PROFILE_VERSION = 1
PROJECT_VERSION = 1
SUPPORTED_IMPORTS = {".pdf", ".docx", ".txt", ".md"}
STOPWORDS = {
    "about", "after", "also", "and", "are", "because", "been", "being", "but", "can", "company",
    "could", "from", "have", "into", "job", "more", "must", "our", "role", "that", "the", "their",
    "them", "they", "this", "through", "using", "what", "when", "where", "which", "will", "with", "work",
    "you", "your", "years", "including", "required", "preferred", "responsibilities", "qualifications",
}
STANDARD_SECTIONS = ("summary", "skills", "experience", "education")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(value: Any, limit: int = 20_000) -> str:
    return " ".join(str(value or "").replace("\x00", " ").split())[:limit]


def _multiline(value: Any, limit: int = 80_000) -> str:
    text = str(value or "").replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.splitlines()).strip()[:limit]


def _slug(value: Any, fallback: str = "resume") -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", _clean(value, 120)).strip("-").lower()
    return slug[:70] or fallback


def _identifier(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def default_profile() -> Dict[str, Any]:
    return {
        "schema_version": PROFILE_VERSION,
        "profile_version": 0,
        "saved": False,
        "contact": {"name": "", "email": "", "phone": "", "location": "", "links": []},
        "target_title": "",
        "headline": "",
        "summary": "",
        "skills": [],
        "career_history": "",
        "education_notes": "",
        "certifications": [],
        "projects": [],
        "awards": [],
        "federal": {
            "citizenship": "",
            "clearance": "",
            "veterans_preference": "",
            "special_hiring_authority": "",
        },
        "sources": [],
        "updated_at": "",
    }


def normalize_profile(value: Dict[str, Any]) -> Dict[str, Any]:
    source = value if isinstance(value, dict) else {}
    profile = default_profile()
    contact = source.get("contact") if isinstance(source.get("contact"), dict) else {}
    profile["contact"] = {
        "name": _clean(contact.get("name"), 200),
        "email": _clean(contact.get("email"), 320),
        "phone": _clean(contact.get("phone"), 100),
        "location": _clean(contact.get("location"), 240),
        "links": [_clean(item, 500) for item in (contact.get("links") or []) if _clean(item, 500)][:8],
    }
    for key in ("target_title", "headline", "summary"):
        profile[key] = _clean(source.get(key), 4000)
    for key in ("career_history", "education_notes"):
        profile[key] = _multiline(source.get(key))
    for key in ("skills", "certifications", "projects", "awards"):
        raw = source.get(key) or []
        if isinstance(raw, str):
            raw = re.split(r"[,\n]", raw)
        profile[key] = [_clean(item, 1000) for item in raw if _clean(item, 1000)][:200]
    federal = source.get("federal") if isinstance(source.get("federal"), dict) else {}
    profile["federal"] = {key: _clean(federal.get(key), 500) for key in profile["federal"]}
    profile["sources"] = [row for row in (source.get("sources") or []) if isinstance(row, dict)][:100]
    profile["profile_version"] = max(0, int(source.get("profile_version") or 0))
    profile["saved"] = bool(source.get("saved"))
    profile["updated_at"] = _clean(source.get("updated_at"), 100)
    return profile


def normalize_draft(value: Dict[str, Any]) -> Dict[str, Any]:
    source = value if isinstance(value, dict) else {}
    contact = source.get("contact") if isinstance(source.get("contact"), dict) else {}
    experiences: List[Dict[str, Any]] = []
    for row in source.get("experience") or []:
        if not isinstance(row, dict):
            continue
        experiences.append({
            "title": _clean(row.get("title"), 300),
            "employer": _clean(row.get("employer"), 300),
            "location": _clean(row.get("location"), 240),
            "start_date": _clean(row.get("start_date"), 80),
            "end_date": _clean(row.get("end_date"), 80),
            "bullets": [_clean(item, 2000) for item in (row.get("bullets") or []) if _clean(item, 2000)][:12],
        })
    education: List[Dict[str, Any]] = []
    for row in source.get("education") or []:
        if isinstance(row, str):
            education.append({"credential": _clean(row, 1000), "school": "", "location": "", "date": ""})
        elif isinstance(row, dict):
            education.append({
                "credential": _clean(row.get("credential"), 600),
                "school": _clean(row.get("school"), 400),
                "location": _clean(row.get("location"), 240),
                "date": _clean(row.get("date"), 80),
            })
    return {
        "contact": {
            "name": _clean(contact.get("name"), 200),
            "email": _clean(contact.get("email"), 320),
            "phone": _clean(contact.get("phone"), 100),
            "location": _clean(contact.get("location"), 240),
            "links": [_clean(item, 500) for item in (contact.get("links") or []) if _clean(item, 500)][:8],
        },
        "target_title": _clean(source.get("target_title"), 300),
        "summary": _clean(source.get("summary"), 4000),
        "skills": [_clean(item, 300) for item in (source.get("skills") or []) if _clean(item, 300)][:80],
        "experience": experiences[:30],
        "education": education[:20],
        "certifications": [_clean(item, 500) for item in (source.get("certifications") or []) if _clean(item, 500)][:50],
        "projects": [_clean(item, 1500) for item in (source.get("projects") or []) if _clean(item, 1500)][:30],
        "awards": [_clean(item, 1000) for item in (source.get("awards") or []) if _clean(item, 1000)][:30],
        "federal_details": [_clean(item, 1200) for item in (source.get("federal_details") or []) if _clean(item, 1200)][:50],
        "cover_letter": _multiline(source.get("cover_letter"), 12_000),
        "linkedin_headline": _clean(source.get("linkedin_headline"), 500),
        "linkedin_about": _multiline(source.get("linkedin_about"), 5000),
        "recruiter_email_subject": _clean(source.get("recruiter_email_subject"), 500),
        "recruiter_email": _multiline(source.get("recruiter_email"), 6000),
        "interview_talking_points": [_clean(item, 2000) for item in (source.get("interview_talking_points") or []) if _clean(item, 2000)][:30],
        "unconfirmed_claims": [
            {
                "claim_id": _clean(row.get("claim_id"), 100) or _identifier("claim"),
                "text": _clean(row.get("text"), 2000),
                "reason": _clean(row.get("reason"), 1000) or "This claim needs your confirmation.",
            }
            for row in (source.get("unconfirmed_claims") or []) if isinstance(row, dict) and _clean(row.get("text"), 2000)
        ],
    }


def extract_json_object(text: str) -> Dict[str, Any]:
    candidate = str(text or "").strip()
    candidate = re.sub(r"^```(?:json)?\s*", "", candidate, flags=re.IGNORECASE)
    candidate = re.sub(r"\s*```$", "", candidate)
    try:
        value = json.loads(candidate)
        if isinstance(value, dict):
            return value
    except json.JSONDecodeError:
        pass
    start, end = candidate.find("{"), candidate.rfind("}")
    if start >= 0 and end > start:
        value = json.loads(candidate[start:end + 1])
        if isinstance(value, dict):
            return value
    raise ValueError("The resume model did not return a valid structured draft.")


def extract_file_text(path: Path) -> str:
    candidate = Path(path).resolve()
    if candidate.suffix.lower() not in SUPPORTED_IMPORTS or not candidate.is_file():
        raise ValueError("Resume imports must be PDF, DOCX, TXT, or Markdown files.")
    if candidate.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support is not installed. Run scripts/setup.ps1.") from exc
        reader = PdfReader(str(candidate))
        text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
    elif candidate.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support is not installed. Run scripts/setup.ps1.") from exc
        document = Document(str(candidate))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
    else:
        text = candidate.read_text(encoding="utf-8", errors="replace").strip()
    if len(text) < 40:
        raise ValueError("The file does not contain enough readable text. Paste the resume text or use a text-based file.")
    return text[:120_000]


def fetch_job_description(url: str) -> str:
    """Fetch bounded public HTTPS job-posting text without allowing local-network access."""
    from urllib.parse import urlparse

    def validate_public_url(candidate: str) -> str:
        checked = _clean(candidate, 2000)
        parsed = urlparse(checked)
        if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
            raise ValueError("Job posting URLs must be public HTTPS links.")
        try:
            addresses = {row[4][0] for row in socket.getaddrinfo(parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM)}
        except OSError as exc:
            raise ValueError("The job posting host could not be resolved.") from exc
        for address in addresses:
            ip = ipaddress.ip_address(address)
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
                raise ValueError("Local or private-network job posting URLs are not allowed.")
        return checked

    value = validate_public_url(url)
    class SafeRedirectHandler(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, request, file_pointer, code, message, headers, new_url):
            return super().redirect_request(request, file_pointer, code, message, headers, validate_public_url(new_url))

    request = urllib.request.Request(value, headers={"User-Agent": "VeridexResumeStudio/1.0"})
    try:
        with urllib.request.build_opener(SafeRedirectHandler()).open(request, timeout=12) as response:
            content_type = str(response.headers.get("Content-Type") or "")
            if "text/html" not in content_type and "text/plain" not in content_type:
                raise ValueError("The job posting URL did not return readable HTML or text.")
            raw = response.read(1_000_001)
            if len(raw) > 1_000_000:
                raise ValueError("The job posting page is too large to import safely.")
            charset = response.headers.get_content_charset() or "utf-8"
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"The job posting could not be retrieved: {exc}") from exc
    page = raw.decode(charset, errors="replace")
    page = re.sub(r"(?is)<(?:script|style|noscript|svg|nav|footer|header)[^>]*>.*?</(?:script|style|noscript|svg|nav|footer|header)>", " ", page)
    page = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h[1-6]>", "\n", page)
    page = re.sub(r"(?s)<[^>]+>", " ", page)
    text = html.unescape(page)
    text = "\n".join(" ".join(line.split()) for line in text.splitlines() if " ".join(line.split()))
    if len(text) < 100:
        raise ValueError("The page did not expose enough readable job-description text. Paste the description instead.")
    return text[:100_000]


def profile_from_text(text: str, source: Dict[str, Any] | None = None) -> Dict[str, Any]:
    raw = _multiline(text, 120_000)
    profile = default_profile()
    lines = [line.strip() for line in raw.splitlines() if line.strip()]
    email = re.search(r"[\w.!#$%&'*+/=?^`{|}~-]+@[\w.-]+\.[A-Za-z]{2,}", raw)
    phone = re.search(r"(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}", raw)
    profile["contact"]["name"] = lines[0][:200] if lines else ""
    profile["contact"]["email"] = email.group(0) if email else ""
    profile["contact"]["phone"] = phone.group(0) if phone else ""
    profile["career_history"] = raw
    profile["sources"] = [source] if source else []
    return profile


def keyword_analysis(job_description: str, resume_text: str = "") -> Dict[str, Any]:
    job = _multiline(job_description, 100_000)
    resume = _multiline(resume_text, 120_000).casefold()
    tokens = re.findall(r"[A-Za-z][A-Za-z0-9+#./-]{2,}", job.casefold())
    counts = Counter(token.strip("./-") for token in tokens if token not in STOPWORDS and not token.isdigit())
    keywords = [word for word, _ in counts.most_common(30)]
    matched = [word for word in keywords if re.search(rf"(?<!\w){re.escape(word)}(?!\w)", resume)]
    missing = [word for word in keywords if word not in matched]
    return {
        "keywords": keywords,
        "matched_keywords": matched,
        "missing_keywords": missing,
        "coverage_percent": round((len(matched) / len(keywords) * 100), 1) if keywords else 0.0,
        "job_description_length": len(job),
    }


def draft_to_text(draft: Dict[str, Any], *, federal: bool = False) -> str:
    row = normalize_draft(draft)
    contact = row["contact"]
    lines = [contact["name"].upper()]
    contact_line = " | ".join(value for value in [contact["location"], contact["phone"], contact["email"], *contact["links"]] if value)
    if contact_line:
        lines.append(contact_line)
    if row["target_title"]:
        lines.extend(["", row["target_title"]])
    if row["summary"]:
        lines.extend(["", "SUMMARY", row["summary"]])
    if row["skills"]:
        lines.extend(["", "SKILLS", " | ".join(row["skills"])])
    if row["experience"]:
        lines.extend(["", "PROFESSIONAL EXPERIENCE"])
        for item in row["experience"]:
            heading = " - ".join(value for value in [item["title"], item["employer"]] if value)
            dates = " - ".join(value for value in [item["start_date"], item["end_date"]] if value)
            lines.append(" | ".join(value for value in [heading, item["location"], dates] if value))
            lines.extend(f"- {bullet}" for bullet in item["bullets"])
    if row["education"]:
        lines.extend(["", "EDUCATION"])
        for item in row["education"]:
            lines.append(" | ".join(value for value in [item["credential"], item["school"], item["location"], item["date"]] if value))
    for title, key in (("CERTIFICATIONS", "certifications"), ("PROJECTS", "projects"), ("AWARDS", "awards")):
        if row[key]:
            lines.extend(["", title, *[f"- {item}" for item in row[key]]])
    if federal and row["federal_details"]:
        lines.extend(["", "FEDERAL APPLICATION DETAILS", *[f"- {item}" for item in row["federal_details"]]])
    return "\n".join(lines).strip() + "\n"


def quality_review(draft: Dict[str, Any], job_description: str = "", *, federal: bool = False) -> Dict[str, Any]:
    row = normalize_draft(draft)
    text = draft_to_text(row, federal=federal)
    lower = text.casefold()
    warnings: List[Dict[str, str]] = []
    if not row["contact"]["name"]:
        warnings.append({"severity": "error", "code": "missing_name", "message": "Add your name."})
    if not row["contact"]["email"]:
        warnings.append({"severity": "error", "code": "missing_email", "message": "Add an email address."})
    for section in STANDARD_SECTIONS:
        if section == "summary" and not row["summary"]:
            warnings.append({"severity": "warning", "code": "missing_summary", "message": "Add a targeted professional summary."})
        elif section == "skills" and not row["skills"]:
            warnings.append({"severity": "warning", "code": "missing_skills", "message": "Add relevant skills from your verified experience."})
        elif section == "experience" and not row["experience"]:
            warnings.append({"severity": "error", "code": "missing_experience", "message": "Add professional or relevant project experience."})
        elif section == "education" and not row["education"]:
            warnings.append({"severity": "warning", "code": "missing_education", "message": "Review whether education or training should be included."})
    bullets = [bullet for item in row["experience"] for bullet in item["bullets"]]
    weak = [bullet for bullet in bullets if re.match(r"^(responsible for|helped|worked on|duties included)\b", bullet, re.I)]
    if weak:
        warnings.append({"severity": "warning", "code": "weak_bullets", "message": f"Rewrite {len(weak)} weak bullet(s) with a clear action and result."})
    if row["unconfirmed_claims"]:
        warnings.append({"severity": "error", "code": "unconfirmed_claims", "message": f"Confirm or remove {len(row['unconfirmed_claims'])} suggested claim(s) before export."})
    if len(text.split()) > (2200 if federal else 1200):
        warnings.append({"severity": "warning", "code": "length", "message": "The resume may be longer than its target format; review pagination."})
    analysis = keyword_analysis(job_description, text) if job_description.strip() else {
        "keywords": [], "matched_keywords": [], "missing_keywords": [], "coverage_percent": 0.0, "job_description_length": 0,
    }
    sections = {name: name.upper() in text for name in ("summary", "skills", "experience", "education")}
    return {
        "ready_to_export": not any(row["severity"] == "error" for row in warnings),
        "warnings": warnings,
        "keyword_analysis": analysis,
        "ats_preview": text,
        "word_count": len(re.findall(r"\b\w+\b", text)),
        "sections": sections,
        "parser_safe": all(token not in lower for token in ("<table", "<textbox")),
    }


class ResumeStudio:
    def __init__(self, data_root: Path, template_root: Path):
        self.data_root = Path(data_root).resolve()
        self.template_root = Path(template_root).resolve()

    def root(self, workspace_id: str) -> Path:
        if not re.fullmatch(r"ws_[a-f0-9]+", str(workspace_id or "")):
            raise KeyError("Unknown workspace")
        return self.data_root / "workspaces" / workspace_id / "resume_studio"

    def profile_path(self, workspace_id: str) -> Path:
        return self.root(workspace_id) / "profile.json"

    def project_path(self, workspace_id: str, project_id: str) -> Path:
        if not re.fullmatch(r"resume_[a-f0-9]+", str(project_id or "")):
            raise KeyError("Unknown resume project")
        return self.root(workspace_id) / "projects" / f"{project_id}.json"

    @staticmethod
    def _write_json(path: Path, value: Any) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(path.suffix + ".tmp")
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def load_profile(self, workspace_id: str) -> Dict[str, Any]:
        path = self.profile_path(workspace_id)
        if not path.is_file():
            return default_profile()
        try:
            profile = normalize_profile(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError, TypeError, ValueError):
            return default_profile()
        profile["saved"] = True
        return profile

    def save_profile(self, workspace_id: str, value: Dict[str, Any], expected_version: int | None = None) -> Dict[str, Any]:
        existing = self.load_profile(workspace_id)
        if expected_version is not None and int(expected_version) != int(existing.get("profile_version") or 0):
            raise ValueError("The career profile changed in another session. Reload it before saving.")
        profile = normalize_profile(value)
        profile["profile_version"] = int(existing.get("profile_version") or 0) + 1
        profile["saved"] = True
        profile["updated_at"] = utc_now()
        self._write_json(self.profile_path(workspace_id), profile)
        return profile

    def list_templates(self) -> List[Dict[str, Any]]:
        rows: List[Dict[str, Any]] = []
        for path in sorted(self.template_root.glob("*.json")):
            try:
                row = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            if isinstance(row, dict) and row.get("id"):
                rows.append(row)
        return rows

    def save_project(self, workspace_id: str, value: Dict[str, Any]) -> Dict[str, Any]:
        source = value if isinstance(value, dict) else {}
        project_id = _clean(source.get("project_id"), 100) or _identifier("resume")
        path = self.project_path(workspace_id, project_id)
        now = utc_now()
        existing: Dict[str, Any] = {}
        if path.is_file():
            try:
                existing = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                existing = {}
        project = {
            "schema_version": PROJECT_VERSION,
            "project_id": project_id,
            "name": _clean(source.get("name"), 300) or "Resume project",
            "resume_type": "federal" if source.get("resume_type") == "federal" else "private",
            "template_id": _clean(source.get("template_id"), 100) or "ats_classic",
            "target_job": {
                "title": _clean((source.get("target_job") or {}).get("title"), 400),
                "company": _clean((source.get("target_job") or {}).get("company"), 400),
                "description": _multiline((source.get("target_job") or {}).get("description"), 100_000),
                "url": _clean((source.get("target_job") or {}).get("url"), 2000),
            },
            "profile_version": int(source.get("profile_version") or 0),
            "draft": normalize_draft(source.get("draft") or {}),
            "review": source.get("review") if isinstance(source.get("review"), dict) else {},
            "source_file_ids": [_clean(item, 100) for item in (source.get("source_file_ids") or []) if _clean(item, 100)],
            "created_at": _clean(existing.get("created_at"), 100) or now,
            "updated_at": now,
        }
        self._write_json(path, project)
        return project

    def list_projects(self, workspace_id: str) -> List[Dict[str, Any]]:
        rows: List[Dict[str, Any]] = []
        root = self.root(workspace_id) / "projects"
        if root.is_dir():
            for path in root.glob("resume_*.json"):
                try:
                    row = json.loads(path.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    continue
                if isinstance(row, dict):
                    rows.append(row)
        return sorted(rows, key=lambda row: str(row.get("updated_at") or ""), reverse=True)


def render_docx(draft: Dict[str, Any], *, federal: bool = False, template: Dict[str, Any] | None = None) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("DOCX support is not installed. Run scripts/setup.ps1.") from exc
    row = normalize_draft(draft)
    style = template or {}
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(float(style.get("margin_top", 0.55)))
    section.bottom_margin = Inches(float(style.get("margin_bottom", 0.55)))
    section.left_margin = Inches(float(style.get("margin_left", 0.62)))
    section.right_margin = Inches(float(style.get("margin_right", 0.62)))
    normal = document.styles["Normal"]
    normal.font.name = str(style.get("body_font", "Arial"))
    normal.font.size = Pt(float(style.get("body_size", 9.5)))
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.04

    def paragraph(text: str = "", *, bold: bool = False, size: float | None = None, align: Any = None, color: str = ""):
        p = document.add_paragraph()
        p.paragraph_format.keep_together = True
        if align is not None:
            p.alignment = align
        run = p.add_run(text)
        run.bold = bold
        run.font.name = str(style.get("body_font", "Arial"))
        run.font.size = Pt(size or float(style.get("body_size", 9.5)))
        if color and re.fullmatch(r"[0-9A-Fa-f]{6}", color):
            run.font.color.rgb = RGBColor.from_string(color.upper())
        return p

    def heading(text: str):
        p = paragraph(text.upper(), bold=True, size=float(style.get("heading_size", 10.5)), color=str(style.get("accent", "1C4D36")))
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        return p

    contact = row["contact"]
    paragraph(contact["name"].upper(), bold=True, size=float(style.get("name_size", 17)), align=WD_ALIGN_PARAGRAPH.CENTER, color=str(style.get("accent", "1C4D36")))
    contact_line = " | ".join(value for value in [contact["location"], contact["phone"], contact["email"], *contact["links"]] if value)
    if contact_line:
        paragraph(contact_line, size=8.7, align=WD_ALIGN_PARAGRAPH.CENTER)
    if row["target_title"]:
        paragraph(row["target_title"], bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    if row["summary"]:
        heading("Summary")
        paragraph(row["summary"])
    if row["skills"]:
        heading("Skills")
        paragraph(" | ".join(row["skills"]))
    if row["experience"]:
        heading("Professional Experience")
        for item in row["experience"]:
            title_line = " - ".join(value for value in [item["title"], item["employer"]] if value)
            dates = " - ".join(value for value in [item["start_date"], item["end_date"]] if value)
            p = paragraph()
            p.paragraph_format.space_before = Pt(4)
            left = p.add_run(title_line)
            left.bold = True
            if item["location"] or dates:
                right = p.add_run(" | " + " | ".join(value for value in [item["location"], dates] if value))
                right.italic = True
            for bullet in item["bullets"]:
                p = document.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.first_line_indent = Inches(-0.14)
                p.paragraph_format.space_after = Pt(1.5)
                p.add_run(bullet)
    if row["education"]:
        heading("Education")
        for item in row["education"]:
            paragraph(" | ".join(value for value in [item["credential"], item["school"], item["location"], item["date"]] if value))
    for title, key in (("Certifications", "certifications"), ("Projects", "projects"), ("Awards", "awards")):
        if row[key]:
            heading(title)
            for item in row[key]:
                p = document.add_paragraph(style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.first_line_indent = Inches(-0.14)
                p.add_run(item)
    if federal and row["federal_details"]:
        heading("Federal Application Details")
        for item in row["federal_details"]:
            p = document.add_paragraph(style="List Bullet")
            p.add_run(item)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def render_pdf(draft: Dict[str, Any], *, federal: bool = False, template: Dict[str, Any] | None = None) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer
        from xml.sax.saxutils import escape
    except ImportError as exc:
        raise RuntimeError("PDF support is not installed. Run scripts/setup.ps1.") from exc
    row = normalize_draft(draft)
    style = template or {}
    stream = io.BytesIO()
    margin_left = float(style.get("margin_left", 0.62)) * inch
    margin_right = float(style.get("margin_right", 0.62)) * inch
    margin_top = float(style.get("margin_top", 0.55)) * inch
    margin_bottom = float(style.get("margin_bottom", 0.55)) * inch
    document = SimpleDocTemplate(stream, pagesize=letter, leftMargin=margin_left, rightMargin=margin_right, topMargin=margin_top, bottomMargin=margin_bottom, title=f"{row['contact']['name']} Resume")
    styles = getSampleStyleSheet()
    accent = colors.HexColor("#" + str(style.get("accent", "1C4D36")))
    body = ParagraphStyle("ResumeBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=float(style.get("body_size", 9.2)), leading=float(style.get("body_size", 9.2)) * 1.24, spaceAfter=3, textColor=colors.HexColor("#1B211D"))
    name_style = ParagraphStyle("ResumeName", parent=body, fontName="Helvetica-Bold", fontSize=float(style.get("name_size", 17)), leading=20, alignment=TA_CENTER, textColor=accent, spaceAfter=3)
    contact_style = ParagraphStyle("ResumeContact", parent=body, fontSize=8.5, leading=10.5, alignment=TA_CENTER, spaceAfter=2)
    heading_style = ParagraphStyle("ResumeHeading", parent=body, fontName="Helvetica-Bold", fontSize=float(style.get("heading_size", 10.5)), leading=13, textColor=accent, spaceBefore=7, spaceAfter=3, borderWidth=0, keepWithNext=True)
    role_style = ParagraphStyle("ResumeRole", parent=body, fontName="Helvetica-Bold", fontSize=9.5, leading=12, spaceBefore=4, spaceAfter=1, keepWithNext=True)
    bullet_style = ParagraphStyle("ResumeBullet", parent=body, leftIndent=12, firstLineIndent=0, bulletIndent=0, spaceAfter=1.5)
    story: List[Any] = []

    def add_heading(value: str) -> None:
        story.append(Paragraph(escape(value.upper()), heading_style))

    contact = row["contact"]
    story.append(Paragraph(escape(contact["name"].upper()), name_style))
    contact_line = " | ".join(value for value in [contact["location"], contact["phone"], contact["email"], *contact["links"]] if value)
    if contact_line:
        story.append(Paragraph(escape(contact_line), contact_style))
    if row["target_title"]:
        target = ParagraphStyle("Target", parent=contact_style, fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceAfter=4)
        story.append(Paragraph(escape(row["target_title"]), target))
    if row["summary"]:
        add_heading("Summary")
        story.append(Paragraph(escape(row["summary"]), body))
    if row["skills"]:
        add_heading("Skills")
        story.append(Paragraph(escape(" | ".join(row["skills"])), body))
    if row["experience"]:
        add_heading("Professional Experience")
        for item in row["experience"]:
            title_line = " - ".join(value for value in [item["title"], item["employer"]] if value)
            meta = " | ".join(value for value in [item["location"], " - ".join(value for value in [item["start_date"], item["end_date"]] if value)] if value)
            block: List[Any] = [Paragraph(escape(title_line + ((" | " + meta) if meta else "")), role_style)]
            if item["bullets"]:
                block.append(ListFlowable([ListItem(Paragraph(escape(bullet), bullet_style)) for bullet in item["bullets"]], bulletType="bullet", start="circle", leftIndent=13, bulletFontSize=5, spaceAfter=2))
            story.append(KeepTogether(block))
    if row["education"]:
        add_heading("Education")
        for item in row["education"]:
            story.append(Paragraph(escape(" | ".join(value for value in [item["credential"], item["school"], item["location"], item["date"]] if value)), body))
    for title, key in (("Certifications", "certifications"), ("Projects", "projects"), ("Awards", "awards")):
        if row[key]:
            add_heading(title)
            story.append(ListFlowable([ListItem(Paragraph(escape(item), bullet_style)) for item in row[key]], bulletType="bullet", start="circle", leftIndent=13, bulletFontName="Helvetica", bulletFontSize=5))
    if federal and row["federal_details"]:
        add_heading("Federal Application Details")
        story.append(ListFlowable([ListItem(Paragraph(escape(item), bullet_style)) for item in row["federal_details"]], bulletType="bullet", start="circle", leftIndent=13, bulletFontName="Helvetica", bulletFontSize=5))

    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#707873"))
        canvas.drawRightString(letter[0] - margin_right, 0.31 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return stream.getvalue()


def render_cover_letter_docx(draft: Dict[str, Any]) -> bytes:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("DOCX support is not installed. Run scripts/setup.ps1.") from exc
    row = normalize_draft(draft)
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(0.75)
        section.left_margin = section.right_margin = Inches(0.8)
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    contact = row["contact"]
    document.add_heading(contact["name"], level=0)
    document.add_paragraph(" | ".join(value for value in [contact["phone"], contact["email"], contact["location"]] if value))
    for block in re.split(r"\n\s*\n", row["cover_letter"]):
        if block.strip():
            document.add_paragraph(block.strip())
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def render_cover_letter_pdf(draft: Dict[str, Any]) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from xml.sax.saxutils import escape
    except ImportError as exc:
        raise RuntimeError("PDF support is not installed. Run scripts/setup.ps1.") from exc
    row = normalize_draft(draft)
    stream = io.BytesIO()
    document = SimpleDocTemplate(stream, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.72 * inch, bottomMargin=0.72 * inch, title=f"{row['contact']['name']} Cover Letter")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("LetterBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=11, textColor=colors.HexColor("#202722"))
    name = ParagraphStyle("LetterName", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor("#1C4D36"), spaceAfter=3)
    contact_style = ParagraphStyle("LetterContact", parent=body, fontSize=8.8, leading=11, textColor=colors.HexColor("#58635C"), spaceAfter=22)
    contact = row["contact"]
    story: List[Any] = [Paragraph(escape(contact["name"]), name)]
    story.append(Paragraph(escape(" | ".join(value for value in [contact["phone"], contact["email"], contact["location"]] if value)), contact_style))
    for block in re.split(r"\n\s*\n", row["cover_letter"]):
        if block.strip():
            story.append(Paragraph(escape(block.strip()).replace("\n", "<br/>"), body))
    document.build(story)
    return stream.getvalue()


def template_by_id(templates: Iterable[Dict[str, Any]], template_id: str) -> Dict[str, Any]:
    return next((row for row in templates if str(row.get("id")) == str(template_id)), {})


def validate_export_bytes(filename: str, content: bytes) -> Dict[str, Any]:
    """Reopen generated output and verify readable content and bounded PDF layout."""
    suffix = Path(filename).suffix.lower()
    if not content:
        raise ValueError(f"Generated {filename} is empty.")
    if suffix == ".txt":
        text = content.decode("utf-8")
        if len(text.strip()) < 40:
            raise ValueError(f"Generated {filename} does not contain enough readable text.")
        return {"format": "text", "text_characters": len(text.strip())}
    if suffix == ".docx":
        try:
            from docx import Document
            document = Document(io.BytesIO(content))
        except Exception as exc:
            raise ValueError(f"Generated {filename} is not a valid DOCX file.") from exc
        text = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
        if len(text) < 40:
            raise ValueError(f"Generated {filename} does not contain enough readable text.")
        return {"format": "docx", "paragraphs": len(document.paragraphs), "text_characters": len(text)}
    if suffix == ".pdf":
        try:
            import fitz
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            extracted = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
            visual = fitz.open(stream=content, filetype="pdf")
        except Exception as exc:
            raise ValueError(f"Generated {filename} is not a valid PDF file.") from exc
        if not reader.pages or len(extracted) < 40:
            raise ValueError(f"Generated {filename} does not contain enough readable text.")
        for page_number, page in enumerate(visual, start=1):
            bounds = page.rect
            for block in page.get_text("blocks"):
                x0, y0, x1, y1 = block[:4]
                if x0 < -1 or y0 < -1 or x1 > bounds.width + 1 or y1 > bounds.height + 1:
                    raise ValueError(f"Generated {filename} has clipped content on page {page_number}.")
        page_count = len(visual)
        visual.close()
        return {"format": "pdf", "pages": page_count, "text_characters": len(extracted), "layout_bounds_valid": True}
    raise ValueError(f"Unsupported generated export format: {suffix}")


def export_bundle(draft: Dict[str, Any], *, resume_type: str, template: Dict[str, Any]) -> Dict[str, bytes]:
    row = normalize_draft(draft)
    if row["unconfirmed_claims"]:
        raise ValueError("Confirm or remove every suggested claim before exporting final files.")
    review = quality_review(row, federal=resume_type == "federal")
    errors = [warning["message"] for warning in review["warnings"] if warning["severity"] == "error"]
    if errors:
        raise ValueError("Export is blocked: " + " ".join(errors))
    federal = resume_type == "federal"
    bundle = {
        "resume.docx": render_docx(row, federal=federal, template=template),
        "resume.pdf": render_pdf(row, federal=federal, template=template),
        "resume.txt": draft_to_text(row, federal=federal).encode("utf-8"),
    }
    if row["cover_letter"]:
        bundle.update({
            "cover-letter.docx": render_cover_letter_docx(row),
            "cover-letter.pdf": render_cover_letter_pdf(row),
            "cover-letter.txt": (row["cover_letter"].strip() + "\n").encode("utf-8"),
        })
    extras: List[str] = []
    if row["linkedin_headline"] or row["linkedin_about"]:
        extras.extend(["LINKEDIN HEADLINE", row["linkedin_headline"], "", "LINKEDIN ABOUT", row["linkedin_about"], ""])
    if row["recruiter_email"]:
        extras.extend(["RECRUITER EMAIL", f"Subject: {row['recruiter_email_subject']}", "", row["recruiter_email"], ""])
    if row["interview_talking_points"]:
        extras.extend(["INTERVIEW TALKING POINTS", *[f"- {item}" for item in row["interview_talking_points"]]])
    if extras:
        bundle["application-kit.txt"] = ("\n".join(extras).strip() + "\n").encode("utf-8")
    for filename, content in bundle.items():
        validate_export_bytes(filename, content)
    return bundle
