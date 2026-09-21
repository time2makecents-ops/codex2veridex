"""Exact-first price-search scope, budget, progress, and cancellation controls.

Provider adapters are deliberately injected. This module decides *when* exact and
similar searches may run, how long they may run, and how results are ordered; it does
not scrape marketplaces or silently broaden the requested scope.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import threading
import time
import uuid
from concurrent.futures import Future, ThreadPoolExecutor, TimeoutError as FutureTimeoutError
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

from openpyxl import load_workbook


EXACT_MATCH_TIERS = ("same_item_candidate", "same_model_or_edition")
MATCH_TIERS = (*EXACT_MATCH_TIERS, "similar", "rejected")
SEARCH_SCOPES = ("exact", "all_likeness")
SEARCH_PRESETS: dict[str, dict[str, int]] = {
    "fast": {"exact_seconds": 90, "all_likeness_seconds": 6 * 60},
    "standard": {"exact_seconds": 3 * 60, "all_likeness_seconds": 12 * 60},
    "extended": {"exact_seconds": 5 * 60, "all_likeness_seconds": 20 * 60},
}

QUERY_SPELLING_CORRECTIONS = {
    "mazine": "magazine",
    "magizine": "magazine",
}
QUERY_TERM_ALIASES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("magazine", ("souvenir program", "program booklet", "publication")),
    ("wine/champagne chiller", ("wine cooler", "champagne cooler")),
    ("wine champagne chiller", ("wine cooler", "champagne cooler")),
    ("blowtorch", ("blowlamp", "gasoline torch")),
    ("blowlamp", ("blowtorch", "gasoline torch")),
    ("t-shirt", ("shirt", "tee")),
    ("textile item", ("t-shirt", "shirt")),
    ("model kit", ("kit", "mechanical construction kit")),
)
QUERY_SOFT_TERMS = frozenset({
    "accessory",
    "collectible",
    "item",
    "magazine",
    "object",
    "publication",
})


def exact_query_variants(query: str, limit: int = 4) -> list[dict[str, Any]]:
    """Build bounded retrieval variants without changing exact-match requirements."""
    maximum = max(1, min(6, int(limit)))
    original = _bounded_text(query, 1000)
    if not original:
        return []

    corrected_tokens = [
        QUERY_SPELLING_CORRECTIONS.get(token.casefold(), token)
        for token in original.split()
    ]
    corrected = " ".join(corrected_tokens)
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()

    def add(value: str, kind: str) -> None:
        normalized = " ".join(str(value or "").split()).strip(" ,/;:-")
        normalized = re.sub(r"(?i)\b([a-z0-9][a-z0-9'-]*)\s*/\s*\1\b", r"\1", normalized)
        identity = normalized.casefold()
        if not normalized or identity in seen or len(rows) >= maximum:
            return
        seen.add(identity)
        rows.append({
            "rank": len(rows) + 1,
            "query": normalized,
            "kind": kind,
        })

    add(corrected, "corrected_full" if corrected.casefold() != original.casefold() else "full")

    slash = re.search(r"(?i)\b([a-z0-9][a-z0-9'-]*)\s*/\s*([a-z0-9][a-z0-9'-]*)\b", corrected)
    if slash:
        for replacement in slash.groups():
            add(corrected[:slash.start()] + replacement + corrected[slash.end():], "alternative_term")

    if re.search(r"(?i)\s+or\s+", corrected):
        add(re.split(r"(?i)\s+or\s+", corrected, maxsplit=1)[0], "uncertain_tail_removed")

    tokens = corrected.split()
    reduced_tokens = [token for token in tokens if token.casefold().strip(".,;:()[]") not in QUERY_SOFT_TERMS]
    if len(reduced_tokens) >= 3 and len(reduced_tokens) < len(tokens):
        add(" ".join(reduced_tokens), "soft_term_removed")

    for source_term, replacements in QUERY_TERM_ALIASES:
        if slash and source_term.casefold() in {value.casefold() for value in slash.groups()}:
            continue
        match = re.search(rf"(?i)\b{re.escape(source_term)}\b", corrected)
        if not match:
            continue
        for replacement in replacements:
            add(corrected[:match.start()] + replacement + corrected[match.end():], "marketplace_alias")

    return rows


def provider_spelling_suggestion(suggestion: Any, original_query: str) -> str:
    """Return a bounded provider spelling suggestion suitable only for retrieval."""
    value = " ".join(str(suggestion or "").split()).strip()
    value = re.sub(r"(?i)^(?:did you mean|showing results for)\s*:?[\s\"'‘’“”]*", "", value)
    value = value.strip(" \"'‘’“”")
    value = re.sub(r"(?i)\s+(?:sold\s+price)\s*$", "", value).strip(" \"'‘’“”")
    value = re.sub(r"(?i)\b([a-z0-9][a-z0-9'-]*)\s*/\s*\1\b", r"\1", value)
    original = " ".join(str(original_query or "").split()).strip()
    if not value or not original or value.casefold() == original.casefold():
        return ""
    original_tokens = set(re.findall(r"[a-z0-9]+", original.casefold()))
    suggested_tokens = set(re.findall(r"[a-z0-9]+", value.casefold()))
    shared = len(original_tokens & suggested_tokens)
    if shared < max(2, min(len(original_tokens), len(suggested_tokens)) // 2):
        return ""
    return value[:1000]


def _utc_now() -> str:
    from datetime import datetime, timezone

    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _bounded_text(value: Any, maximum: int = 1000) -> str:
    return " ".join(str(value or "").split())[:maximum]


@dataclass(frozen=True)
class SearchRequest:
    item_id: str
    query: str
    scope: str = "exact"
    preset: str = "standard"

    @classmethod
    def from_value(cls, value: Mapping[str, Any]) -> "SearchRequest":
        item_id = _bounded_text(value.get("item_id") or value.get("case_id") or "item", 160)
        query = _bounded_text(value.get("query"), 1000)
        scope = _bounded_text(value.get("scope") or "exact", 40).lower().replace("-", "_")
        preset = _bounded_text(value.get("preset") or "standard", 40).lower()
        if not query:
            raise ValueError("price search requires a non-empty query")
        if scope not in SEARCH_SCOPES:
            raise ValueError("scope must be exact or all_likeness")
        if preset not in SEARCH_PRESETS:
            raise ValueError("preset must be fast, standard, or extended")
        return cls(item_id=item_id, query=query, scope=scope, preset=preset)

    @property
    def similar_approved(self) -> bool:
        return self.scope == "all_likeness"

    @property
    def budget_seconds(self) -> int:
        values = SEARCH_PRESETS[self.preset]
        return values["exact_seconds"] if self.scope == "exact" else values["all_likeness_seconds"]

    @property
    def exact_stage_seconds(self) -> int:
        return SEARCH_PRESETS[self.preset]["exact_seconds"]


class SearchCancelled(RuntimeError):
    pass


class SearchTimedOut(RuntimeError):
    pass


@dataclass(frozen=True)
class SearchPhaseContext:
    phase: str
    request: SearchRequest
    phase_started: float
    phase_deadline: float
    overall_deadline: float
    canceled: Callable[[], bool]
    clock: Callable[[], float] = time.monotonic

    @property
    def remaining_seconds(self) -> float:
        return max(0.0, min(self.phase_deadline, self.overall_deadline) - self.clock())

    def checkpoint(self) -> None:
        if self.canceled():
            raise SearchCancelled("Price search canceled")
        if self.remaining_seconds <= 0:
            raise SearchTimedOut(f"{self.phase.replace('_', ' ').title()} search budget expired")


SearchAdapter = Callable[[SearchRequest, SearchPhaseContext], Mapping[str, Any] | Sequence[Mapping[str, Any]]]
ProgressCallback = Callable[[int, str], None]
ResultFinalizer = Callable[[dict[str, Any]], dict[str, Any]]


def _candidate_identity(row: Mapping[str, Any]) -> str:
    for key in ("source_listing_key", "listing_or_lot_id", "source_url", "observation_id"):
        value = _bounded_text(row.get(key), 1000)
        if value:
            return f"{key}:{value.casefold()}"
    stable = json.dumps(
        {
            "title": _bounded_text(row.get("title"), 500).casefold(),
            "platform": _bounded_text(row.get("platform"), 120).casefold(),
            "amount": row.get("amount"),
            "sold_at": _bounded_text(row.get("sold_at"), 80),
        },
        sort_keys=True,
        separators=(",", ":"),
    )
    return "fallback:" + hashlib.sha256(stable.encode("utf-8")).hexdigest()


def _adapter_payload(value: Mapping[str, Any] | Sequence[Mapping[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    if isinstance(value, Mapping):
        raw_candidates = value.get("candidates") or value.get("observations") or []
        raw_sources = value.get("sources") or []
        metadata = {key: item for key, item in value.items() if key not in {"candidates", "observations", "sources"}}
    else:
        raw_candidates = value
        raw_sources = []
        metadata = {}
    candidates = [dict(row) for row in raw_candidates if isinstance(row, Mapping)]
    sources = [dict(row) for row in raw_sources if isinstance(row, Mapping)]
    for row in candidates:
        tier = _bounded_text(row.get("match_tier"), 80)
        if tier not in MATCH_TIERS:
            row["match_tier"] = "rejected"
            differences = list(row.get("differences") or []) if isinstance(row.get("differences"), list) else []
            row["differences"] = [*differences, "Adapter returned an unsupported match tier"]
    return candidates, sources, metadata


def _ordered_unique(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    tier_rank = {tier: index for index, tier in enumerate(MATCH_TIERS)}
    unique: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for row in rows:
        identity = _candidate_identity(row)
        prior = unique.get(identity)
        if prior is None:
            unique[identity] = row
            order.append(identity)
            continue
        if tier_rank.get(str(row.get("match_tier")), 99) < tier_rank.get(str(prior.get("match_tier")), 99):
            unique[identity] = row
    values = [unique[identity] for identity in order]
    values.sort(key=lambda row: tier_rank.get(str(row.get("match_tier")), 99))
    return values


class PriceSearchController:
    def run(
        self,
        value: SearchRequest | Mapping[str, Any],
        exact_search: SearchAdapter,
        similar_search: SearchAdapter | None = None,
        *,
        progress: ProgressCallback | None = None,
        canceled: Callable[[], bool] | None = None,
        monotonic: Callable[[], float] = time.monotonic,
    ) -> dict[str, Any]:
        request = value if isinstance(value, SearchRequest) else SearchRequest.from_value(value)
        progress = progress or (lambda _value, _message: None)
        canceled = canceled or (lambda: False)
        started = monotonic()
        overall_deadline = started + request.budget_seconds
        exact_deadline = min(overall_deadline, started + request.exact_stage_seconds)
        phase_records: list[dict[str, Any]] = []
        all_candidates: list[dict[str, Any]] = []
        all_sources: list[dict[str, Any]] = []

        def run_phase(name: str, adapter: SearchAdapter, deadline: float, start_progress: int, message: str) -> None:
            phase_started = monotonic()
            context = SearchPhaseContext(name, request, phase_started, deadline, overall_deadline, canceled, monotonic)
            context.checkpoint()
            progress(start_progress, message)
            returned = adapter(request, context)
            candidates, sources, metadata = _adapter_payload(returned)
            all_candidates.extend(candidates)
            all_sources.extend(sources)
            phase_record = {
                "phase": name,
                "status": "completed",
                "candidate_count": len(candidates),
                "source_count": len(sources),
                "elapsed_seconds": round(monotonic() - phase_started, 3),
                "metadata": metadata,
            }
            phase_records.append(phase_record)
            try:
                context.checkpoint()
            except SearchTimedOut:
                phase_record["status"] = "timed_out"
                raise

        progress(3, "Preparing exact-match search")
        try:
            run_phase("exact", exact_search, exact_deadline, 10, "Searching exact likeness")
            exact_after_first = [row for row in _ordered_unique(all_candidates) if row.get("match_tier") in EXACT_MATCH_TIERS]
            progress(58, f"Exact search complete: {len(exact_after_first)} match{'es' if len(exact_after_first) != 1 else ''}")
            if request.scope == "all_likeness":
                if similar_search is None:
                    raise ValueError("all_likeness scope requires a similar-search adapter")
                run_phase("similar", similar_search, overall_deadline, 64, "Searching similar likeness after exact results")
                progress(94, "Ordering exact results before similar results")
        except SearchCancelled:
            raise
        except SearchTimedOut as exc:
            ordered = _ordered_unique(all_candidates)
            exact = [row for row in ordered if row.get("match_tier") in EXACT_MATCH_TIERS]
            similar = [row for row in ordered if row.get("match_tier") == "similar"] if request.similar_approved else []
            return self._result(
                request, "timed_out", exact, similar,
                [row for row in ordered if row.get("match_tier") == "rejected"],
                all_sources, phase_records, started, monotonic(), str(exc),
            )

        ordered = _ordered_unique(all_candidates)
        exact = [row for row in ordered if row.get("match_tier") in EXACT_MATCH_TIERS]
        similar = [row for row in ordered if row.get("match_tier") == "similar"] if request.similar_approved else []
        rejected = [row for row in ordered if row.get("match_tier") == "rejected"]
        status = "completed" if exact or request.similar_approved else "similar_approval_required"
        progress(100, "Price search complete" if status == "completed" else "Exact search complete; similar search needs approval")
        return self._result(request, status, exact, similar, rejected, all_sources, phase_records, started, monotonic(), "")

    @staticmethod
    def _result(
        request: SearchRequest,
        status: str,
        exact: list[dict[str, Any]],
        similar: list[dict[str, Any]],
        rejected: list[dict[str, Any]],
        sources: list[dict[str, Any]],
        phases: list[dict[str, Any]],
        started: float,
        finished: float,
        message: str,
    ) -> dict[str, Any]:
        return {
            "schema": "veridex.price_search.result.v1",
            "status": status,
            "item_id": request.item_id,
            "query": request.query,
            "scope": request.scope,
            "preset": request.preset,
            "budget_seconds": request.budget_seconds,
            "elapsed_seconds": round(max(0.0, finished - started), 3),
            "similar_approved": request.similar_approved,
            "exact_results": exact,
            "similar_results": similar,
            "ordered_results": [*exact, *similar],
            "rejected_results": rejected,
            "sources": sources,
            "phases": phases,
            "message": message,
            "similar_search_prompt": "No exact match found. Search similar items?" if not exact and not request.similar_approved else "",
        }


class PriceSearchJobManager:
    def __init__(self, workers: int = 2):
        self._jobs: dict[str, dict[str, Any]] = {}
        self._futures: dict[str, Future[Any]] = {}
        self._cancelled: set[str] = set()
        self._lock = threading.RLock()
        self._executor = ThreadPoolExecutor(max_workers=max(1, int(workers)), thread_name_prefix="veridex-price-search")
        self._controller = PriceSearchController()

    def submit(
        self,
        value: SearchRequest | Mapping[str, Any],
        exact_search: SearchAdapter,
        similar_search: SearchAdapter | None = None,
        *,
        finalize: ResultFinalizer | None = None,
    ) -> dict[str, Any]:
        request = value if isinstance(value, SearchRequest) else SearchRequest.from_value(value)
        job_id = f"pricejob_{uuid.uuid4().hex[:16]}"
        now = _utc_now()
        with self._lock:
            self._jobs[job_id] = {
                "job_id": job_id,
                "status": "queued",
                "progress": 0,
                "message": "Queued",
                "created_at": now,
                "updated_at": now,
                "item_id": request.item_id,
                "query": request.query,
                "scope": request.scope,
                "preset": request.preset,
                "budget_seconds": request.budget_seconds,
            }

        def progress(value: int, message: str) -> None:
            with self._lock:
                self._jobs[job_id].update({
                    "progress": max(0, min(100, int(value))),
                    "message": _bounded_text(message, 500),
                    "updated_at": _utc_now(),
                })

        def canceled() -> bool:
            with self._lock:
                return job_id in self._cancelled

        def work() -> None:
            with self._lock:
                self._jobs[job_id].update({"status": "running", "updated_at": _utc_now()})
            try:
                result = self._controller.run(
                    request,
                    exact_search,
                    similar_search,
                    progress=progress,
                    canceled=canceled,
                )
                if finalize is not None and not canceled():
                    progress(100, "Saving governed price evidence")
                    result = finalize(result)
                with self._lock:
                    if canceled():
                        self._jobs[job_id].update({"status": "canceled", "progress": 100, "message": "Canceled", "updated_at": _utc_now()})
                    else:
                        self._jobs[job_id].update({
                            "status": "completed",
                            "progress": 100,
                            "message": "Complete" if result["status"] == "completed" else result["status"].replace("_", " ").title(),
                            "result": result,
                            "updated_at": _utc_now(),
                        })
            except SearchCancelled:
                with self._lock:
                    self._jobs[job_id].update({"status": "canceled", "progress": 100, "message": "Canceled", "updated_at": _utc_now()})
            except Exception as exc:
                with self._lock:
                    self._jobs[job_id].update({"status": "failed", "progress": 100, "message": str(exc), "error": str(exc), "updated_at": _utc_now()})

        future = self._executor.submit(work)
        with self._lock:
            self._futures[job_id] = future
        return self.get(job_id)

    def get(self, job_id: str) -> dict[str, Any]:
        with self._lock:
            if job_id not in self._jobs:
                raise KeyError("Unknown price-search job")
            return dict(self._jobs[job_id])

    def wait(self, job_id: str, timeout: float | None = None) -> dict[str, Any]:
        with self._lock:
            if job_id not in self._jobs:
                raise KeyError("Unknown price-search job")
            future = self._futures.get(job_id)
        if future is not None:
            try:
                future.result(timeout=timeout)
            except FutureTimeoutError:
                pass
        return self.get(job_id)

    def cancel(self, job_id: str) -> dict[str, Any]:
        with self._lock:
            if job_id not in self._jobs:
                raise KeyError("Unknown price-search job")
            if self._jobs[job_id]["status"] in {"completed", "failed", "canceled"}:
                return dict(self._jobs[job_id])
            self._cancelled.add(job_id)
            self._jobs[job_id].update({"status": "canceling", "message": "Canceling", "updated_at": _utc_now()})
            return dict(self._jobs[job_id])


@dataclass(frozen=True)
class WorkbookSearchSubject:
    item_id: str
    row_number: int
    description: str
    maker: str
    title_or_model: str
    query: str


def workbook_search_subjects(workbook_path: Path) -> dict[str, Any]:
    """Read search subjects from Inventory without modifying the workbook."""
    workbook_path = Path(workbook_path).resolve()
    if not workbook_path.is_file() or workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("supply an existing XLSX or XLSM inventory workbook")
    source_hash = _sha256(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if "Inventory" not in workbook.sheetnames:
            raise ValueError("workbook is missing the Inventory sheet")
        sheet = workbook["Inventory"]
        headings = {
            _bounded_text(cell.value, 160).casefold(): cell.column
            for cell in sheet[1]
            if _bounded_text(cell.value, 160)
        }
        required = ("description",)
        if any(value not in headings for value in required):
            raise ValueError("Inventory sheet is missing the Description column")
        item_column = headings.get("item #")
        maker_column = headings.get("maker / brand")
        title_column = headings.get("title / model")
        subjects = []
        for row_number in range(2, sheet.max_row + 1):
            description = _bounded_text(sheet.cell(row_number, headings["description"]).value, 500)
            maker = _bounded_text(sheet.cell(row_number, maker_column).value, 300) if maker_column else ""
            title = _bounded_text(sheet.cell(row_number, title_column).value, 300) if title_column else ""
            if not any((description, maker, title)):
                continue
            item_value = sheet.cell(row_number, item_column).value if item_column else row_number - 1
            item_id = _bounded_text(item_value, 120) or str(row_number - 1)
            query = _bounded_text(" ".join(value for value in (maker, title, description) if value), 1000)
            subjects.append(WorkbookSearchSubject(item_id, row_number, description, maker, title, query))
    finally:
        workbook.close()
    if _sha256(workbook_path) != source_hash:
        raise RuntimeError("source workbook changed while search subjects were read")
    return {
        "schema": "veridex.price_search.workbook_subjects.v1",
        "source_path": str(workbook_path),
        "source_sha256": source_hash,
        "subject_count": len(subjects),
        "subjects": [asdict(subject) for subject in subjects],
    }
