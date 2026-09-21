"""Batch item crops and inventory spreadsheet generation.

The module is intentionally UI-agnostic.  Browser canvas selections are stored in
original-image coordinates, then rendered here as rectangular context crops and
lasso-isolated PNGs.  Source images are never modified.
"""

from __future__ import annotations

import hashlib
import io
import json
import textwrap
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFilter, ImageOps


INVENTORY_COLUMNS = (
    "Item #",
    "Image",
    "Description",
    "Maker / Brand",
    "Title / Model",
    "Visible Asking Price",
    "Prior Provisional Estimate Low",
    "Prior Provisional Estimate High",
    "Currency",
    "Exact Match",
    "Similar Match",
    "Comparable Sales",
    "Source Links",
    "Link Status",
    "Confidence",
    "Review Status",
    "Unreadable Reason",
    "Notes",
)

PRICE_SUMMARY_COLUMNS = (
    "Pricing Status",
    "Exact Sold Count",
    "Expected Resale Low",
    "Expected Resale Median",
    "Expected Resale High",
    "Last Sold Price",
    "Last Sold Date",
    "Last Sold Site",
    "Last Sold Venue",
    "Match Tier Used",
    "Matching Engagement",
    "Pricing Confidence",
    "Research Date",
)

ENGAGEMENT_FIELDS = ("watchers", "likes", "favorites", "saves", "bids", "unique_bidders", "views", "quantity_sold")

PRICE_EVIDENCE_COLUMNS = (
    "Item #",
    "Match Tier",
    "Match Confidence",
    "Valuation Eligible",
    "Exclusion Reasons",
    "Sale Status",
    "Price Basis",
    "Amount",
    "Valuation Amount",
    "Currency",
    "Shipping",
    "Buyer Premium",
    "Sold Date",
    "Platform",
    "Venue",
    "Listing / Lot ID",
    "Title",
    "Source URL",
    "Evidence Status",
    "Watchers",
    "Likes",
    "Favorites",
    "Saves",
    "Bids",
    "Unique Bidders",
    "Views",
    "Quantity Sold",
    "Engagement Captured At",
    "Evidence Captured At",
)


@dataclass(frozen=True)
class ItemSelection:
    item_number: int
    points: tuple[tuple[int, int], ...]
    description: str = ""
    maker: str = ""
    title: str = ""
    visible_asking_price: str = ""
    confidence: str = ""
    review_status: str = "Pending research"
    unreadable_reason: str = ""
    notes: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if self.item_number < 1:
            raise ValueError("item_number must be positive")
        if len(self.points) < 2:
            raise ValueError("a selection needs at least two points")


def _bounded_points(points: Sequence[tuple[int, int]], width: int, height: int) -> list[tuple[int, int]]:
    return [
        (max(0, min(width - 1, int(x))), max(0, min(height - 1, int(y))))
        for x, y in points
    ]


def _selection_box(points: Sequence[tuple[int, int]], width: int, height: int, padding: int) -> tuple[int, int, int, int]:
    bounded = _bounded_points(points, width, height)
    left = max(0, min(point[0] for point in bounded) - padding)
    top = max(0, min(point[1] for point in bounded) - padding)
    right = min(width, max(point[0] for point in bounded) + padding + 1)
    bottom = min(height, max(point[1] for point in bounded) + padding + 1)
    if right <= left or bottom <= top:
        raise ValueError("selection has no visible area")
    return left, top, right, bottom


def render_selection(
    source: Image.Image,
    selection: ItemSelection,
    *,
    padding: int = 12,
    neutral_background: tuple[int, int, int] = (245, 245, 245),
) -> dict[str, Image.Image]:
    """Return context, transparent, and neutral-background images for a selection."""
    image = ImageOps.exif_transpose(source).convert("RGBA")
    points = _bounded_points(selection.points, image.width, image.height)
    box = _selection_box(points, image.width, image.height, max(0, int(padding)))
    context = image.crop(box)

    mask = Image.new("L", image.size, 0)
    draw = ImageDraw.Draw(mask)
    if len(points) == 2:
        draw.rectangle((points[0], points[1]), fill=255)
    else:
        draw.polygon(points, fill=255)
    local_mask = mask.crop(box)

    transparent = Image.new("RGBA", context.size, (0, 0, 0, 0))
    transparent.paste(context, (0, 0), local_mask)
    neutral = Image.new("RGB", context.size, neutral_background)
    neutral.paste(context.convert("RGB"), (0, 0), local_mask)
    return {"context": context.convert("RGB"), "transparent": transparent, "neutral": neutral}


def export_item_images(
    source_path: Path,
    selections: Iterable[ItemSelection],
    output_dir: Path,
    *,
    padding: int = 12,
) -> list[dict[str, Any]]:
    """Export derived images without altering the source and return row metadata."""
    source_path = Path(source_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    with Image.open(source_path) as opened:
        source = ImageOps.exif_transpose(opened).copy()

    rows: list[dict[str, Any]] = []
    for selection in selections:
        rendered = render_selection(source, selection, padding=padding)
        stem = f"item_{selection.item_number:03d}"
        context_path = output_dir / f"{stem}_context.jpg"
        transparent_path = output_dir / f"{stem}_isolated.png"
        neutral_path = output_dir / f"{stem}_neutral.jpg"
        rendered["context"].save(context_path, format="JPEG", quality=94, optimize=True)
        rendered["transparent"].save(transparent_path, format="PNG", optimize=True)
        rendered["neutral"].save(neutral_path, format="JPEG", quality=94, optimize=True)
        rows.append(
            {
                "selection": selection,
                "context_path": context_path,
                "transparent_path": transparent_path,
                "neutral_path": neutral_path,
            }
        )
    return rows


def _embed_thumbnail(sheet: Any, cell: str, image_path: Path) -> None:
    thumbnail = SpreadsheetImage(str(image_path))
    thumbnail.width = 120
    thumbnail.height = 90
    sheet.add_image(thumbnail, cell)


def create_inventory_workbook(
    output_path: Path,
    source_image: Path,
    items: Sequence[dict[str, Any]],
) -> Path:
    """Create a reviewable XLSX with embedded crops and a needs-review sheet."""
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    inventory = workbook.active
    inventory.title = "Inventory"
    inventory.append(INVENTORY_COLUMNS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in inventory[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    review = workbook.create_sheet("Needs Review")
    review.append(("Item #", "Crop", "Reason", "Requested Follow-up"))
    for cell in review[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="9C0006")

    for row_index, item in enumerate(items, start=2):
        selection: ItemSelection = item["selection"]
        values = (
            selection.item_number,
            "",
            selection.description,
            selection.maker,
            selection.title,
            selection.visible_asking_price,
            "",
            "",
            "USD",
            "Pending",
            "Not searched",
            "",
            "",
            "Unchecked",
            selection.confidence,
            selection.review_status,
            selection.unreadable_reason,
            selection.notes,
        )
        inventory.append(values)
        inventory.row_dimensions[row_index].height = 72
        _embed_thumbnail(inventory, f"B{row_index}", Path(item["context_path"]))
        for column in (13,):
            cell = inventory.cell(row=row_index, column=column)
            if cell.value:
                cell.hyperlink = str(cell.value).splitlines()[0]
                cell.style = "Hyperlink"

        if selection.unreadable_reason or selection.review_status.lower().startswith("needs"):
            review_row = review.max_row + 1
            review.append(
                (
                    selection.item_number,
                    "",
                    selection.unreadable_reason or "Identification needs confirmation",
                    "Provide a closer, straighter photo or adjust the selection.",
                )
            )
            review.row_dimensions[review_row].height = 92
            _embed_thumbnail(review, f"B{review_row}", Path(item["context_path"]))

    inventory.freeze_panes = "A2"
    inventory.auto_filter.ref = f"A1:R{max(1, inventory.max_row)}"
    widths = [10, 20, 32, 20, 24, 18, 18, 18, 10, 18, 18, 26, 38, 14, 12, 18, 28, 36]
    for index, width in enumerate(widths, start=1):
        inventory.column_dimensions[get_column_letter(index)].width = width
    for row in inventory.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    review.freeze_panes = "A2"
    for column, width in enumerate((10, 20, 40, 48), start=1):
        review.column_dimensions[get_column_letter(column)].width = width
    for row in review.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    info = workbook.create_sheet("Source & Instructions")
    info.append(("Source image", str(Path(source_image).resolve())))
    info.append(("Search order", "Original context crop first; isolated/neutral versions only as fallback."))
    info.append(("Exact-match policy", "Do not mix similar items into exact-match values or comparable-sale calculations."))
    info.append(("Unreadable policy", "Keep the item row, explain the limitation, and include its crop on Needs Review."))
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(output_path)
    check = load_workbook(output_path, read_only=True, data_only=False)
    required = {"Inventory", "Needs Review", "Source & Instructions"}
    if not required.issubset(check.sheetnames):
        check.close()
        raise ValueError("generated workbook is missing required sheets")
    check.close()
    return output_path


def _workbook_image_hashes(workbook_path: Path) -> list[tuple[str, int, int, str]]:
    """Return stable embedded-image identities without relying on media filenames."""
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    hashes: list[tuple[str, int, int, str]] = []
    try:
        for sheet in workbook.worksheets:
            for image in getattr(sheet, "_images", []):
                anchor = getattr(image, "anchor", None)
                marker = getattr(anchor, "_from", None)
                row = int(getattr(marker, "row", -1))
                column = int(getattr(marker, "col", -1))
                hashes.append((sheet.title, row, column, hashlib.sha256(image._data()).hexdigest()))
    finally:
        workbook.close()
    return hashes


def _price_status(evaluation: dict[str, Any]) -> str:
    expected = evaluation.get("expected_resale") if isinstance(evaluation.get("expected_resale"), dict) else {}
    summary = evaluation.get("evidence_summary") if isinstance(evaluation.get("evidence_summary"), dict) else {}
    if int(expected.get("sold_count") or 0) > 0:
        return "Exact sold evidence"
    if int(summary.get("exact") or len(evaluation.get("exact_results") or [])) > 0:
        return "Exact match found; no supported sold price"
    return "No exact match"


def _matching_engagement(evaluation: dict[str, Any]) -> str:
    snapshots = []
    for observation in evaluation.get("exact_results") or []:
        if not isinstance(observation, dict):
            continue
        engagement = observation.get("engagement") if isinstance(observation.get("engagement"), dict) else {}
        values = [
            f"{engagement[field]} {field.replace('_', ' ')}"
            for field in ENGAGEMENT_FIELDS
            if field in engagement and engagement[field] is not None
        ]
        if not values:
            continue
        source = str(observation.get("platform") or observation.get("source_id") or "Source")
        captured = str(observation.get("engagement_captured_at") or observation.get("captured_at") or "")
        snapshots.append(f"{source}: {', '.join(values)}" + (f" ({captured})" if captured else ""))
    return "\n".join(snapshots)


def apply_price_evaluations_to_workbook(
    source_workbook: Path,
    output_path: Path,
    evaluations: dict[Any, dict[str, Any]],
    *,
    evidence_json_path: Path | None = None,
) -> Path:
    """Create a priced workbook copy while preserving rows and embedded item images."""
    source_workbook = Path(source_workbook).resolve()
    output_path = Path(output_path).resolve()
    if source_workbook == output_path:
        raise ValueError("priced workbook output must not overwrite the source workbook")
    if not source_workbook.is_file():
        raise FileNotFoundError(source_workbook)

    source_file_hash = hashlib.sha256(source_workbook.read_bytes()).hexdigest()
    source_image_hashes = _workbook_image_hashes(source_workbook)
    workbook = load_workbook(source_workbook, read_only=False, data_only=False)
    if "Inventory" not in workbook.sheetnames:
        workbook.close()
        raise ValueError("source workbook is missing the Inventory sheet")
    inventory = workbook["Inventory"]
    original_inventory_rows = inventory.max_row

    header_map = {
        str(cell.value or "").strip(): cell.column
        for cell in inventory[1]
        if str(cell.value or "").strip()
    }
    legacy_headers = {
        "Estimated Value Low": "Prior Provisional Estimate Low",
        "Estimated Value High": "Prior Provisional Estimate High",
    }
    for old, new in legacy_headers.items():
        if old in header_map:
            inventory.cell(1, header_map[old]).value = new
    header_map = {
        str(cell.value or "").strip(): cell.column
        for cell in inventory[1]
        if str(cell.value or "").strip()
    }
    for heading in PRICE_SUMMARY_COLUMNS:
        if heading not in header_map:
            column = inventory.max_column + 1
            cell = inventory.cell(1, column, heading)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            inventory.column_dimensions[get_column_letter(column)].width = 20
            header_map[heading] = column

    normalized_evaluations: dict[str, dict[str, Any]] = {}
    for item_number, value in evaluations.items():
        if not isinstance(value, dict):
            continue
        evaluation = value.get("price_evaluation") if isinstance(value.get("price_evaluation"), dict) else value
        normalized_evaluations[str(item_number)] = evaluation

    item_column = header_map.get("Item #", 1)
    for row_index in range(2, inventory.max_row + 1):
        item_number = inventory.cell(row_index, item_column).value
        evaluation = normalized_evaluations.get(str(item_number))
        if evaluation is None:
            continue
        expected = evaluation.get("expected_resale") if isinstance(evaluation.get("expected_resale"), dict) else {}
        last_sold = evaluation.get("last_sold") if isinstance(evaluation.get("last_sold"), dict) else {}
        captured_dates = [
            str(row.get("captured_at") or "")
            for row in evaluation.get("observations") or []
            if isinstance(row, dict) and row.get("captured_at")
        ]
        values = {
            "Pricing Status": _price_status(evaluation),
            "Exact Sold Count": expected.get("sold_count"),
            "Expected Resale Low": expected.get("low"),
            "Expected Resale Median": expected.get("median"),
            "Expected Resale High": expected.get("high"),
            "Last Sold Price": last_sold.get("amount"),
            "Last Sold Date": last_sold.get("sold_at"),
            "Last Sold Site": last_sold.get("platform") or last_sold.get("source_id"),
            "Last Sold Venue": last_sold.get("venue"),
            "Match Tier Used": expected.get("match_tier"),
            "Matching Engagement": _matching_engagement(evaluation),
            "Pricing Confidence": expected.get("confidence"),
            "Research Date": max(captured_dates, default=""),
        }
        for heading, value in values.items():
            cell = inventory.cell(row_index, header_map[heading], value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if "Price Evidence" in workbook.sheetnames:
        workbook.remove(workbook["Price Evidence"])
    evidence_sheet = workbook.create_sheet("Price Evidence")
    evidence_sheet.append(PRICE_EVIDENCE_COLUMNS)
    for cell in evidence_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for item_number, evaluation in normalized_evaluations.items():
        for observation in evaluation.get("observations") or []:
            if not isinstance(observation, dict):
                continue
            engagement = observation.get("engagement") if isinstance(observation.get("engagement"), dict) else {}
            evidence_sheet.append((
                item_number,
                observation.get("match_tier"),
                observation.get("match_confidence"),
                observation.get("valuation_eligible"),
                ", ".join(observation.get("valuation_exclusion_reasons") or []),
                observation.get("sale_status"),
                observation.get("price_basis"),
                observation.get("amount"),
                observation.get("valuation_amount"),
                observation.get("currency"),
                observation.get("shipping"),
                observation.get("buyer_premium"),
                observation.get("sold_at"),
                observation.get("platform") or observation.get("source_id"),
                observation.get("venue"),
                observation.get("listing_or_lot_id"),
                observation.get("title"),
                observation.get("source_url"),
                observation.get("evidence_status"),
                *(engagement.get(field) if field in engagement else None for field in ENGAGEMENT_FIELDS),
                observation.get("engagement_captured_at"),
                observation.get("captured_at"),
            ))
            source_cell = evidence_sheet.cell(evidence_sheet.max_row, PRICE_EVIDENCE_COLUMNS.index("Source URL") + 1)
            if source_cell.value:
                source_cell.hyperlink = str(source_cell.value)
                source_cell.style = "Hyperlink"

    evidence_sheet.freeze_panes = "A2"
    evidence_sheet.auto_filter.ref = f"A1:{get_column_letter(evidence_sheet.max_column)}{max(1, evidence_sheet.max_row)}"
    for column in range(1, evidence_sheet.max_column + 1):
        evidence_sheet.column_dimensions[get_column_letter(column)].width = 18
    for row in evidence_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if "Source & Instructions" in workbook.sheetnames:
        info = workbook["Source & Instructions"]
    else:
        info = workbook.create_sheet("Source & Instructions")
    existing_labels = {str(info.cell(row, 1).value or "") for row in range(1, info.max_row + 1)}
    policies = (
        ("Pricing evidence", "Expected resale uses only supported, medium/high-confidence exact sold observations. Active asking prices and similar items are excluded."),
        ("Engagement policy", "Watchers, likes, favorites, saves, bids, views, and quantity sold are timestamped snapshots and do not affect valuation."),
        ("Missing counts", "A blank engagement field means unavailable; zero is recorded only when the source explicitly reports zero."),
    )
    for label, explanation in policies:
        if label not in existing_labels:
            info.append((label, explanation))
            info.cell(info.max_row, 1).font = Font(bold=True)
            info.cell(info.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    inventory.auto_filter.ref = f"A1:{get_column_letter(inventory.max_column)}{max(1, inventory.max_row)}"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    check = load_workbook(output_path, read_only=True, data_only=False)
    try:
        if check["Inventory"].max_row != original_inventory_rows:
            raise ValueError("priced workbook changed the inventory row count")
        if "Price Evidence" not in check.sheetnames:
            raise ValueError("priced workbook is missing the Price Evidence sheet")
    finally:
        check.close()
    if not output_path.is_file() or output_path.stat().st_size <= 0:
        raise ValueError("priced workbook was not written")
    if _workbook_image_hashes(output_path) != source_image_hashes:
        raise ValueError("priced workbook did not preserve embedded inventory images")
    if hashlib.sha256(source_workbook.read_bytes()).hexdigest() != source_file_hash:
        raise ValueError("source workbook changed while creating the priced copy")

    if evidence_json_path is not None:
        evidence_json_path = Path(evidence_json_path).resolve()
        evidence_json_path.parent.mkdir(parents=True, exist_ok=True)
        evidence_json_path.write_text(
            json.dumps({
                "schema": "veridex.price_evidence.v1",
                "source_workbook": str(source_workbook),
                "source_sha256": source_file_hash,
                "evaluations": normalized_evaluations,
            }, ensure_ascii=False, indent=2, sort_keys=True, default=str),
            encoding="utf-8",
        )
    return output_path


def create_review_contact_sheet(
    output_path: Path,
    items: Sequence[dict[str, Any]],
    *,
    columns: int = 2,
) -> Path:
    """Create a labeled PNG containing every crop that still needs human review."""
    review_items = [
        item
        for item in items
        if item["selection"].unreadable_reason
        or item["selection"].review_status.lower().startswith("needs")
    ]
    if not review_items:
        raise ValueError("no review items were supplied")
    columns = max(1, min(4, int(columns)))
    tile_width, image_height, text_height = 520, 330, 120
    tile_height = image_height + text_height
    rows = (len(review_items) + columns - 1) // columns
    sheet = Image.new("RGB", (tile_width * columns, tile_height * rows), "white")
    draw = ImageDraw.Draw(sheet)
    for index, item in enumerate(review_items):
        selection: ItemSelection = item["selection"]
        x = (index % columns) * tile_width
        y = (index // columns) * tile_height
        with Image.open(item["context_path"]) as opened:
            crop = ImageOps.contain(ImageOps.exif_transpose(opened).convert("RGB"), (tile_width - 24, image_height - 24))
        image_x = x + (tile_width - crop.width) // 2
        image_y = y + (image_height - crop.height) // 2
        sheet.paste(crop, (image_x, image_y))
        draw.rectangle((x, y, x + tile_width - 1, y + tile_height - 1), outline=(120, 120, 120), width=2)
        heading = f"Item {selection.item_number}: {selection.description or 'Unidentified item'}"
        reason = selection.unreadable_reason or "Identification needs confirmation."
        draw.text((x + 12, y + image_height + 8), "\n".join(textwrap.wrap(heading, width=58)), fill="black")
        draw.text((x + 12, y + image_height + 50), "\n".join(textwrap.wrap(reason, width=68)), fill=(140, 0, 0))
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # Contact sheets favor fast review generation over smaller file size.
    sheet.save(output_path, format="PNG", compress_level=3)
    return output_path


def create_annotated_selection_image(
    source_path: Path,
    selections: Sequence[ItemSelection],
    output_path: Path,
    *,
    scale: int = 3,
) -> Path:
    """Render numbered selection boundaries over an enlarged copy of the source."""
    scale = max(1, min(8, int(scale)))
    with Image.open(source_path) as opened:
        source = ImageOps.exif_transpose(opened).convert("RGB")
    canvas = source.resize((source.width * scale, source.height * scale), Image.Resampling.NEAREST)
    draw = ImageDraw.Draw(canvas)
    colors = ((255, 40, 40), (30, 110, 255), (30, 180, 80), (255, 140, 0), (180, 40, 220))
    line_width = max(2, scale)
    for index, selection in enumerate(selections):
        points = _bounded_points(selection.points, source.width, source.height)
        scaled = [(x * scale, y * scale) for x, y in points]
        color = colors[index % len(colors)]
        if len(scaled) == 2:
            left = min(scaled[0][0], scaled[1][0])
            top = min(scaled[0][1], scaled[1][1])
            right = max(scaled[0][0], scaled[1][0])
            bottom = max(scaled[0][1], scaled[1][1])
            draw.rectangle((left, top, right, bottom), outline=color, width=line_width)
            label_x, label_y = left, top
        else:
            draw.line((*scaled, scaled[0]), fill=color, width=line_width, joint="curve")
            label_x = min(point[0] for point in scaled)
            label_y = min(point[1] for point in scaled)
        label = str(selection.item_number)
        label_box = draw.textbbox((label_x, label_y), label)
        pad = max(2, scale)
        draw.rectangle(
            (label_box[0] - pad, label_box[1] - pad, label_box[2] + pad, label_box[3] + pad),
            fill=color,
        )
        draw.text((label_x, label_y), label, fill="white", stroke_width=1, stroke_fill="black")
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path, format="PNG", compress_level=3)
    return output_path


def create_layered_sharpen_enhancement(
    source_path: Path,
    output_path: Path,
    *,
    scale: float = 1.4,
    sharpen_opacity: float = 0.4,
    radius: float = 1.2,
    percent: int = 170,
    threshold: int = 2,
) -> Path:
    """Enlarge, sharpen a duplicate layer, and blend it over the enlarged base.

    This is a deterministic viewing aid. It does not recover or infer missing
    source detail and must not be treated as independent visual evidence.
    """
    scale = max(1.0, min(8.0, float(scale)))
    sharpen_opacity = max(0.0, min(1.0, float(sharpen_opacity)))
    with Image.open(source_path) as opened:
        source = ImageOps.exif_transpose(opened).convert("RGB")
    width = max(1, round(source.width * scale))
    height = max(1, round(source.height * scale))
    base = source.resize((width, height), Image.Resampling.LANCZOS)
    sharpened = base.filter(
        ImageFilter.UnsharpMask(
            radius=max(0.1, float(radius)),
            percent=max(1, int(percent)),
            threshold=max(0, int(threshold)),
        )
    )
    composite = Image.blend(base, sharpened, sharpen_opacity)
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    composite.save(output_path, format="PNG", compress_level=3)
    return output_path


def create_enhancement_variants(
    source_path: Path,
    output_dir: Path,
    *,
    scale: float = 4.0,
) -> list[dict[str, Any]]:
    """Create deterministic viewing and OCR variants of one item crop.

    The variants expose only transformations of source pixels. Thresholded
    variants are explicitly marked as OCR-only because their altered strokes
    are unsuitable for visual-identification evidence.
    """
    source_path = Path(source_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    scale = max(1.0, min(8.0, float(scale)))

    with Image.open(source_path) as opened:
        source = ImageOps.exif_transpose(opened).convert("RGB")
    size = (max(1, round(source.width * scale)), max(1, round(source.height * scale)))
    enlarged = source.resize(size, Image.Resampling.LANCZOS)
    rgb = np.asarray(enlarged)
    bgr = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)

    lab = cv2.cvtColor(bgr, cv2.COLOR_BGR2LAB)
    luminance, channel_a, channel_b = cv2.split(lab)
    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
    clahe_bgr = cv2.cvtColor(
        cv2.merge((clahe.apply(luminance), channel_a, channel_b)),
        cv2.COLOR_LAB2BGR,
    )

    denoised = cv2.fastNlMeansDenoisingColored(bgr, None, 3, 3, 7, 21)
    blurred = cv2.GaussianBlur(denoised, (0, 0), 1.0)
    denoised_sharpened = cv2.addWeighted(denoised, 1.45, blurred, -0.45, 0)

    gray = cv2.cvtColor(clahe_bgr, cv2.COLOR_BGR2GRAY)
    minimum_dimension = min(gray.shape)
    block_size = min(31, minimum_dimension if minimum_dimension % 2 == 1 else minimum_dimension - 1)
    block_size = max(3, block_size)
    adaptive = cv2.adaptiveThreshold(
        gray,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        block_size,
        7,
    )
    _, otsu = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)

    stem = source_path.stem
    specifications = (
        ("original", source, "Evidence", "Unaltered source crop"),
        ("lanczos_4x", enlarged, "Viewing / search", f"Lanczos enlargement at {scale:g}x"),
        (
            "clahe_color",
            Image.fromarray(cv2.cvtColor(clahe_bgr, cv2.COLOR_BGR2RGB)),
            "Viewing / search",
            "Local contrast enhancement on luminance",
        ),
        (
            "denoise_sharpen",
            Image.fromarray(cv2.cvtColor(denoised_sharpened, cv2.COLOR_BGR2RGB)),
            "Viewing / search",
            "Mild color denoise followed by unsharp masking",
        ),
        ("adaptive_threshold", Image.fromarray(adaptive), "OCR only", "Adaptive Gaussian threshold"),
        ("otsu_threshold", Image.fromarray(otsu), "OCR only", "Otsu threshold after CLAHE"),
    )
    variants: list[dict[str, Any]] = []
    for key, image, use, description in specifications:
        output_path = output_dir / f"{stem}_{key}.png"
        image.save(output_path, format="PNG", compress_level=3)
        variants.append(
            {
                "key": key,
                "path": output_path,
                "use": use,
                "description": description,
                "source_path": source_path,
            }
        )
    return variants


def create_enhancement_comparison_sheet(
    output_path: Path,
    item_variants: Sequence[tuple[str, Sequence[dict[str, Any]]]],
    *,
    tile_width: int = 260,
    image_height: int = 190,
) -> Path:
    """Create a labeled grid with one item per row and one variant per column."""
    if not item_variants:
        raise ValueError("no enhancement variants were supplied")
    expected_keys = tuple(variant["key"] for variant in item_variants[0][1])
    if not expected_keys:
        raise ValueError("each item needs at least one enhancement variant")
    for _, variants in item_variants:
        if tuple(variant["key"] for variant in variants) != expected_keys:
            raise ValueError("every item must contain the same ordered variants")

    header_height, label_height = 64, 46
    row_height = image_height + label_height
    sheet = Image.new(
        "RGB",
        (tile_width * len(expected_keys), header_height + row_height * len(item_variants)),
        "white",
    )
    draw = ImageDraw.Draw(sheet)
    for column, variant in enumerate(item_variants[0][1]):
        x = column * tile_width
        heading = f"{variant['key'].replace('_', ' ').title()}\n{variant['use']}"
        draw.multiline_text((x + 8, 8), heading, fill="black", spacing=3)
        draw.rectangle((x, 0, x + tile_width - 1, header_height - 1), outline=(100, 100, 100), width=2)

    for row, (item_label, variants) in enumerate(item_variants):
        y = header_height + row * row_height
        for column, variant in enumerate(variants):
            x = column * tile_width
            with Image.open(variant["path"]) as opened:
                preview = ImageOps.contain(
                    ImageOps.exif_transpose(opened).convert("RGB"),
                    (tile_width - 16, image_height - 12),
                )
            preview_x = x + (tile_width - preview.width) // 2
            preview_y = y + (image_height - preview.height) // 2
            sheet.paste(preview, (preview_x, preview_y))
            draw.rectangle((x, y, x + tile_width - 1, y + row_height - 1), outline=(150, 150, 150), width=1)
            draw.text((x + 8, y + image_height + 5), item_label, fill="black")
            if variant["use"] == "OCR only":
                draw.text((x + 8, y + image_height + 22), "ALTERED - OCR ONLY", fill=(170, 0, 0))

    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output_path, format="PNG", compress_level=3)
    return output_path
